import hashlib
import io
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

from runner_agent.authoring import Recording, draft_workbook
from runner_agent.repair import check_candidate, export_repair, require_supported_step


@pytest.fixture
def source(tmp_path):
    recording = Recording()
    recording.accept({"action": "fill", "locator": "input:nth-of-type(1)"})
    recording.accept({"action": "click", "locator": "button:nth-of-type(1)"})
    wb = load_workbook(io.BytesIO(draft_workbook(recording)))
    for sheet in (wb["steps"], wb["testcases"]):
        column = list(next(sheet.values)).index("active") + 1
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row, column, "Y")
    wb.create_sheet("custom_notes").append(["Keep this sheet", "synthetic note"])
    path = tmp_path / "original.xlsx"
    wb.save(path)
    wb.close()
    return path


def test_export_preserves_all_sheets_and_data_but_deactivates_copy(source, tmp_path):
    original = source.read_bytes()
    output = tmp_path / "repaired.xlsx"
    result = export_repair(source, output, 2, "input:nth-of-type(2)", hashlib.sha256(original).hexdigest())
    before, after = load_workbook(io.BytesIO(original)), load_workbook(output)
    assert source.read_bytes() == original
    assert result["draft_sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    for sheet in before:
        old_rows, new_rows = list(sheet.values), list(after[sheet.title].values)
        headers = old_rows[0]
        for row_index, (old, new) in enumerate(zip(old_rows, new_rows)):
            for column, (old_value, new_value) in enumerate(zip(old, new)):
                if row_index and sheet.title in {"steps", "testcases"} and headers[column] == "active":
                    assert new_value == "N"
                elif row_index == 1 and sheet.title == "steps" and headers[column] == "locator":
                    assert new_value == "input:nth-of-type(2)"
                else:
                    assert old_value == new_value
    assert "repair_review" in after.sheetnames
    before.close()
    after.close()


def test_export_refuses_stale_workbook_and_existing_output(source, tmp_path):
    old_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    source.write_bytes(source.read_bytes() + b"changed")
    output = tmp_path / "repaired.xlsx"
    with pytest.raises(ValueError, match="changed"):
        export_repair(source, output, 2, "input:nth-of-type(2)", old_hash)
    assert not output.exists()
    output.write_bytes(b"existing")
    with pytest.raises(ValueError):
        export_repair(source, output, 2, "input:nth-of-type(2)", old_hash)
    assert output.read_bytes() == b"existing"


@pytest.mark.parametrize("step", [{"action": "wait"}, {"action": "radio"}, {"action": "form_item"},
                                   {"action": "fill", "group": "repeat"},
                                   {"action": "read_result_single", "read_method": "label_input"}])
def test_complex_semantics_require_manual_repair(step):
    with pytest.raises(ValueError):
        require_supported_step(step)


@pytest.mark.parametrize("count,visible,identity,expected", [(1, True, True, "READY_FOR_REVIEW"),
    (0, True, True, "NOT_UNIQUE"), (2, True, True, "NOT_UNIQUE"), (1, False, True, "HIDDEN"),
    (1, True, False, "STALE_SELECTION")])
def test_candidate_must_still_be_the_selected_visible_element(count, visible, identity, expected):
    locator = SimpleNamespace(count=lambda: count, is_visible=lambda: visible, evaluate=lambda script: identity)
    page = SimpleNamespace(locator=lambda selector: locator)
    assert check_candidate(page, {"action": "fill"}, "input:nth-of-type(1)") == expected
    assert check_candidate(page, {"action": "fill"}, "button:nth-of-type(1)") == "INCOMPATIBLE_TAG"


@pytest.mark.parametrize("confirmed,stale", [(False, False), (True, False), (True, True)])
def test_cli_exports_only_after_confirmation(source, tmp_path, monkeypatch, confirmed, stale):
    import playwright.sync_api
    from repair_runner import repair
    closed = []
    page = SimpleNamespace(main_frame=object())
    live = {"identity": True}
    locator = SimpleNamespace(count=lambda: 1, is_visible=lambda: True, evaluate=lambda script: live["identity"])
    page.locator = lambda selector: locator
    class Context:
        pages = [page]
        def expose_binding(self, name, callback): self.callback = callback
        def add_init_script(self, script): assert "runnerRepairPick" in script
        def new_page(self):
            self.callback({"page": page, "frame": page.main_frame}, {"locator": "input:nth-of-type(2)"})
            return page
    context = Context()
    browser = SimpleNamespace(new_context=lambda: context, is_connected=lambda: True, close=lambda: closed.append(True))
    class Playwright:
        def __enter__(self): return SimpleNamespace(chromium=SimpleNamespace(launch=lambda **kw: browser))
        def __exit__(self, *args): pass
    monkeypatch.setattr(playwright.sync_api, "sync_playwright", Playwright)
    output = tmp_path / "reviewed.xlsx"
    def answer(_):
        if stale:
            live["identity"] = False
            context.pages = []
        return "EXPORT" if confirmed else "q"
    result = repair(source, 2, output, prompt=answer)
    assert output.exists() == (confirmed and not stale)
    assert bool(result) == (confirmed and not stale)
    assert closed == [True]
