import io

from openpyxl import load_workbook
import pytest

from compose_runner import compose
from runner_agent.authoring import _workbook
from runner_agent.preflight import check
from runner_agent.preparation import merge
from src.runner.discovery import Snapshot, DiscoveryPlan


def draft(screen, name, action="fill"):
    snapshot = Snapshot.model_validate({"candidates": [
        {"id": "c1", "kind": "input", "selector": "input:nth-of-type(1)"}]})
    step = {"screen": screen, "step": name, "action": action,
            "value_source": "testcase" if action == "fill" else "empty", "candidate_id": "c1"}
    if action == "read_result_single":
        step.update(read_method="css_input", match_type="exact")
    return _workbook(DiscoveryPlan.model_validate({"steps": [step]}).bind(snapshot))


def test_multi_screen_discovery_prepare_and_preflight_without_generating_user_data(tmp_path):
    first, second, combined = [tmp_path / name for name in ("first.xlsx", "second.xlsx", "flow.xlsx")]
    first.write_bytes(draft("search", "customer_ref"))
    second.write_bytes(draft("result", "read_amount", "read_result_single"))
    before = [first.read_bytes(), second.read_bytes()]
    assert compose([first, second], combined) == {"drafts": 2, "steps": 2, "screens": 2}
    wb = load_workbook(combined)
    assert wb["testcases"].max_row == 1 and "settings" not in wb.sheetnames
    assert list(wb["testcases"].values)[0][-2:] == ("customer_ref", "expected_amount")
    assert [r[0] for r in list(wb["steps"].values)[1:]] == ["search", "result"]
    # Simulate an existing USER workbook, then merge without approving any settings changes.
    settings = wb.create_sheet("settings")
    for row in [("key", "value"), ("url", "https://example.test"), ("screen_flow", "search,result"),
                ("result_screen", "result"), ("default_timeout", "4321")]:
        settings.append(row)
    wb["testcases"].append(["USER_CASE", "synthetic user data", "N", "RM", "X", "FAKE_REF", "123"])
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    prepared = merge(buffer.getvalue(), combined.read_bytes())
    report = check(prepared)
    assert not report["ready_for_local_review"]
    wb = load_workbook(io.BytesIO(prepared))
    assert dict(list(wb["settings"].values)[1:])["default_timeout"] == "4321"
    assert list(wb["testcases"].values)[1][-2:] == ("FAKE_REF", "123")
    # Only the human explicitly activates their cases/steps.
    for row in range(2, wb["steps"].max_row + 1):
        wb["steps"].cell(row, 7, "Y")
    wb["testcases"].cell(2, 3, "Y")
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    assert check(buffer.getvalue())["ready_for_local_review"]
    assert [first.read_bytes(), second.read_bytes()] == before
    with pytest.raises(ValueError):
        compose([first, second], combined)


@pytest.mark.parametrize("problem", ["duplicate", "active", "testcase", "settings", "screen_order", "result_screens", "column_collision"])
def test_compose_rejects_conflicts_without_writing_or_changing_sources(tmp_path, problem):
    first, second, third = [tmp_path / name for name in ("first.xlsx", "second.xlsx", "third.xlsx")]
    first.write_bytes(draft("one", "field_one"))
    second.write_bytes(draft("two", "field_two"))
    drafts = [first, second]
    if problem == "duplicate":
        second.write_bytes(draft("two", "field_one"))
    elif problem == "screen_order":
        third.write_bytes(draft("one", "field_three"))
        drafts.append(third)
    elif problem == "result_screens":
        first.write_bytes(draft("one", "read_first", "read_result_single"))
        second.write_bytes(draft("two", "read_second", "read_result_single"))
    elif problem == "column_collision":
        first.write_bytes(draft("one", "expected_amount"))
        second.write_bytes(draft("two", "read_amount", "read_result_single"))
    else:
        wb = load_workbook(second)
        if problem == "active":
            wb["steps"].cell(2, 7, "Y")
        if problem == "testcase":
            wb["testcases"].append(["USER_ENTERED", "Keep this"])
        if problem == "settings":
            wb.create_sheet("settings").append(["key", "value"])
        wb.save(second)
        wb.close()
    originals = [p.read_bytes() for p in drafts]
    output = tmp_path / "output.xlsx"
    with pytest.raises(ValueError):
        compose(drafts, output)
    assert not output.exists()
    assert [p.read_bytes() for p in drafts] == originals
