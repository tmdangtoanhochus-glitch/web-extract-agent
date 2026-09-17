import io
import json

import httpx
import pytest
from fastapi import FastAPI
from fastapi.responses import JSONResponse
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api.runner import create_runner_router
from src.runner.planner import Plan, PlanError, StepPlanner, validate_description
from src.runner.repository import Repository
from src.runner.service import RunnerError, Service
from runner_agent.authoring import planned_workbook

DESCRIPTION = "Điền tài khoản role RM rồi bấm đăng nhập."
VALID = {
    "steps": [
        {"screen": "login", "step": "username", "action": "fill", "value_source": "account"},
        {"screen": "login", "step": "password", "action": "fill", "value_source": "account"},
        {"screen": "login", "step": "submit", "action": "click", "value_source": "empty"}],
}


def planner_with(content, status=200, requests=None):
    def handle(request):
        if requests is not None:
            requests.append(request)
        return httpx.Response(status, json={"choices": [{"message": {"content": content}}]})
    client = httpx.Client(transport=httpx.MockTransport(handle))
    return StepPlanner("https://ai.example/v1", "synthetic-key", "synthetic-model", client=client)


def test_chat_contract_and_inactive_workbook():
    requests = []
    plan = planner_with(json.dumps(VALID), requests=requests).plan(DESCRIPTION)
    assert str(requests[0].url) == "https://ai.example/v1/chat/completions"
    payload = json.loads(requests[0].content)
    assert payload["model"] == "synthetic-model"
    assert payload["messages"][1]["content"] == DESCRIPTION
    wb = load_workbook(io.BytesIO(planned_workbook(plan)))
    assert all(row[4] == ":not(*)" and row[6] == "N" for row in list(wb["steps"].values)[1:])
    assert wb["testcases"].max_row == 1
    assert "password" in [row[1] for row in list(wb["steps"].values)[1:]]
    assert len(list(wb["steps"].values)[0]) == 13
    assert "settings" not in wb.sheetnames
    wb.close()


@pytest.mark.parametrize("content", [
    "not json", '{"steps": []}',
    json.dumps({"steps": [{"action": "exec", "target": "button"}]}),
    json.dumps({"steps": [{"action": "click", "target": "password"}]}),
    json.dumps({"steps": [{"action": "fill", "target": "username", "value": "synthetic-private"}]}),
    json.dumps({**VALID, "code": "synthetic-private"}),
    "x" * 32001,
    json.dumps({"steps": VALID["steps"] * 34}),
])
def test_rejects_untrusted_model_output_without_echo(content):
    with pytest.raises(PlanError) as error:
        planner_with(content).plan(DESCRIPTION)
    assert "synthetic-private" not in str(error.value)


@pytest.mark.parametrize("text", ["password: synthetic-private", "OTP=123456", "Bearer synthetic-private",
                                  "Mở https://internal.example", "ngắn"])
def test_sensitive_description_rejected_before_network(text):
    requests = []
    with pytest.raises(PlanError):
        planner_with(json.dumps(VALID), requests=requests).plan(text)
    assert requests == []


def test_placeholder_description_is_allowed():
    assert validate_description("Điền password: ${RM_PASSWORD} rồi đăng nhập.")


def test_http_failure_is_generic(caplog):
    with pytest.raises(PlanError):
        planner_with("synthetic-private", status=500).plan(DESCRIPTION)
    assert "synthetic-private" not in caplog.text


@pytest.fixture
def api(tmp_path):
    repo = Repository()
    service = Service(repo, tmp_path / "server")
    owner = service.add_user("tester", "synthetic-password-123")
    token = service.login("tester", "synthetic-password-123")["session"]
    def build(planner):
        app = FastAPI()
        @app.exception_handler(RunnerError)
        async def handler(req, exc):
            return JSONResponse(status_code=exc.code, content={"detail": str(exc)})
        app.include_router(create_runner_router(service, planner))
        return TestClient(app)
    yield build, repo, {"Authorization": "Bearer " + token}, owner
    repo.close()


def test_api_requires_login_consent_and_opt_in(api):
    build, _, auth, _ = api
    client = build(None)
    body = {"description": DESCRIPTION, "reviewed_no_secrets": True}
    assert client.post("/runner/authoring/describe", json=body).status_code == 401
    assert client.get("/runner/authoring/capabilities", headers=auth).json()["describe"] is False
    assert client.post("/runner/authoring/describe", headers=auth, json=body).status_code == 503
    client = build(planner_with(json.dumps(VALID)))
    assert client.post("/runner/authoring/describe", headers=auth, json={
        **body, "reviewed_no_secrets": False}).status_code == 422


def test_api_streams_draft_without_saving_content_or_creating_run(api):
    build, repo, auth, owner = api
    client = build(planner_with(json.dumps(VALID)))
    response = client.post("/runner/authoring/describe", headers=auth, json={
        "description": DESCRIPTION, "reviewed_no_secrets": True})
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    wb = load_workbook(io.BytesIO(response.content))
    assert "steps" in wb.sheetnames
    wb.close()
    with repo.transaction():
        assert repo.all("runs") == []
        audits = repo.all("audit")
    generated = [a for a in audits if a["event"] == "DRAFT_GENERATED"]
    assert generated[0]["user_id"] == owner["id"]
    assert DESCRIPTION not in json.dumps(audits, ensure_ascii=False)
