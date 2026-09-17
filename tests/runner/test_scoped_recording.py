import io
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

from record_runner import receive_event
from runner_agent.authoring import Recording, draft_workbook
from runner_agent.inspector import runner_locator_builder
from runner_agent.upload import local_upload_path
from src.runner.scoped_locator import PREFIX, decode, encode, resolve


def test_recorder_javascript_with_synthetic_dom():
    import subprocess
    from playwright._impl._driver import compute_driver_executable
    node, _ = compute_driver_executable()
    result = subprocess.run([str(node), str(Path(__file__).with_name("recorder_dom.cjs"))],
                            capture_output=True, text=True, timeout=20,
                            env={"SYSTEMROOT": "C:\\Windows"})
    assert result.returncode == 0, result.stderr
    assert "checks passed" in result.stdout


def scoped(tag="input"):
    return encode([{"kind": "shadow", "css": "app-root:nth-of-type(1)"},
                   {"kind": "frame", "css": "iframe:nth-of-type(1)"},
                   {"kind": "shadow", "css": "form-panel:nth-of-type(1)"},
                   {"kind": "target", "css": tag + ":nth-of-type(1)"}])


class BrowserDouble:
    def __init__(self): self.calls = []
    def locator(self, css): self.calls.append(("locator", css)); return self
    def frame_locator(self, css): self.calls.append(("frame", css)); return self
    def click(self): self.calls.append(("click",))
    def uncheck(self): self.calls.append(("uncheck",))
    def check(self): self.calls.append(("check",))
    def set_input_files(self, path): self.calls.append(("upload", path))
    def wait_for(self, **kw): self.calls.append(("wait", kw))
    def input_value(self): return "synthetic"


def test_nested_frame_shadow_paths_resolve_in_order():
    page = BrowserDouble()
    assert resolve(page, scoped()) is page
    assert page.calls == [("locator", "app-root:nth-of-type(1)"), ("frame", "iframe:nth-of-type(1)"),
                          ("locator", "form-panel:nth-of-type(1)"), ("locator", "input:nth-of-type(1)")]


def test_typing_coalesces_only_between_other_frame_actions():
    recording = Recording()
    fill = {"action": "fill", "locator": "input:nth-of-type(1)"}
    assert recording.accept(fill)
    assert recording.accept(fill)
    assert len(recording.events) == 1
    assert recording.accept({"action": "click", "locator": scoped("button")})
    assert recording.accept(fill)
    assert [e["action"] for e in recording.events] == ["fill", "click", "fill"]
    assert recording.accept({"action": "fill", "locator": scoped()})
    assert len(recording.events) == 4


def test_pause_and_screen_changes_end_a_typing_sequence():
    recording = Recording()
    fill = {"action": "fill", "locator": "input:nth-of-type(1)"}
    assert recording.accept(fill)
    assert recording.control({"control": "toggle_pause"})
    assert not recording.accept(fill)
    assert recording.control({"control": "toggle_pause"})
    assert recording.accept(fill)
    assert recording.control({"control": "next_screen"})
    assert recording.accept(fill)
    assert [e["screen"] for e in recording.events] == ["recorded", "recorded", "recorded_002"]
    assert recording.suppressed == 1 and recording.dropped == 0


def test_typing_at_event_limit_does_not_drop_keystrokes_or_other_actions_silently():
    recording = Recording(limit=1)
    fill = {"action": "fill", "locator": "input:nth-of-type(1)"}
    for _ in range(20):
        assert recording.accept(fill)
    assert len(recording.events) == 1 and recording.dropped == 0
    assert not recording.accept({"action": "click", "locator": scoped("button")})
    assert not recording.accept(fill)
    assert recording.dropped == 2


@pytest.mark.parametrize("value", [PREFIX + "{}", PREFIX + "[]", PREFIX + "null",
    PREFIX + '[{"kind":"target","css":"[value=private]"}]',
    PREFIX + '[{"kind":"frame","css":"button:nth-of-type(1)"},{"kind":"target","css":"input:nth-of-type(1)"}]',
    PREFIX + '[{"kind":"target","css":"input:nth-of-type(1)","value":"private"}]'])
def test_scope_rejects_data_and_invalid_structure(value):
    with pytest.raises((ValueError, TypeError)):
        decode(value)
    assert not Recording().accept({"action": "fill", "locator": value})


def test_extended_recording_contains_only_actions_and_empty_headers():
    recording = Recording()
    for action in ("check", "uncheck", "upload", "read_result_single"):
        assert recording.accept({"action": action, "locator": scoped()})
    wb = load_workbook(io.BytesIO(draft_workbook(recording)))
    try:
        rows = list(wb["steps"].values)[1:]
        assert [row[2] for row in rows] == ["check", "uncheck", "upload", "read_result_single"]
        assert all(row[6] == "N" for row in rows)
        assert rows[2][5] == "testcase"
        assert wb["testcases"].max_row == 1 and "settings" not in wb.sheetnames
        assert list(next(wb["testcases"].values))[-2:] == ["field_0003", "expected_field_0004"]
    finally:
        wb.close()
    assert not recording.accept({"action": "upload", "locator": "button:nth-of-type(1)"})


def test_binding_prepends_iframe_ancestry_without_reading_url_or_name():
    disposed = []
    main = object()
    class Frame:
        parent_frame = main
        def frame_element(self):
            return SimpleNamespace(evaluate=lambda script: [{"kind": "target", "css": "iframe:nth-of-type(2)"}],
                                   dispose=lambda: disposed.append(True))
    recording = Recording()
    state = receive_event(recording, {"frame": Frame(), "page": SimpleNamespace(main_frame=main)},
                          {"action": "fill", "locator": "input:nth-of-type(3)"}, "synthetic function")
    assert state["accepted"] and disposed == [True]
    assert decode(recording.events[0]["locator"]) == [
        {"kind": "frame", "css": "iframe:nth-of-type(2)"}, {"kind": "target", "css": "input:nth-of-type(3)"}]


def test_detached_iframe_rejected_without_exporting_raw_error():
    class Frame:
        parent_frame = object()
        def frame_element(self): raise RuntimeError("synthetic-private")
    recording = Recording()
    state = receive_event(recording, {"frame": Frame(), "page": SimpleNamespace(main_frame=object())},
                          {"action": "click", "locator": "button:nth-of-type(1)"}, "function")
    assert not state["accepted"] and recording.dropped == 1 and not recording.events
    assert "synthetic-private" not in str(state)


@pytest.mark.parametrize("name", [".env", "runner.env", "secrets/file.txt", "credentials.json", "private.pem"])
def test_upload_blocks_protected_paths_before_file_access(name, monkeypatch):
    monkeypatch.setattr(Path, "is_file", lambda *_: pytest.fail("must not inspect protected file"))
    with pytest.raises(ValueError, match="Protected"):
        local_upload_path(name)


def test_executor_resolves_scope_for_click_wait_read_and_new_actions(tmp_path):
    build = runner_locator_builder()
    runtime = build.__globals__
    page = BrowserDouble()
    assert build(page, "css", scoped()) is page
    runtime["do_click"](page, "css", scoped("button"), scoped(), "test")
    assert sum(call == ("click",) for call in page.calls) == 1
    assert page.calls[-1] == ("wait", {"state": "visible"})
    assert runtime["read_result_field"](page, {"read_method": "css_input", "locator": scoped()}, 0) == "synthetic"
    path = tmp_path / "synthetic.txt"
    path.write_text("synthetic content")
    step = {"action": "upload", "locator_type": "css", "locator": scoped(),
            "wait_selector": "", "step": "attachment"}
    runtime["run_step"](page, step, str(path))
    assert page.calls[-1] == ("upload", str(path))
    runtime["run_step"](page, {**step, "action": "uncheck"}, "")
    assert page.calls[-1] == ("uncheck",)
