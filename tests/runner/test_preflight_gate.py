import io
import json
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from openpyxl import Workbook

from runner_agent import executor
from runner_agent.client import LocalAgent
from runner_agent.preflight import check, VALID_ACTIONS, VALID_LOCATOR_TYPES
from src.runner.preflight_contract import PreflightReport, metadata
from src.runner.workbook_schema import STEP_COLUMNS, BASE_CASE_COLUMNS


def workbook(active="Y", locator="button", action="click"):
    wb = Workbook()
    wb.active.title = "settings"
    for row in [("key", "value"), ("url", "https://example.test"), ("screen_flow", "home")]:
        wb.active.append(row)
    steps = wb.create_sheet("steps")
    steps.append(STEP_COLUMNS)
    data = {"screen": "home", "step": "submit", "action": action, "locator_type": "css",
            "locator": locator, "value_source": "empty", "active": active}
    steps.append([data.get(key, "") for key in STEP_COLUMNS])
    cases = wb.create_sheet("testcases")
    cases.append(BASE_CASE_COLUMNS)
    cases.append(["TC1", "synthetic-private-description", active, "RM", "X"])
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


@pytest.mark.parametrize("content,code", [(workbook(active="N"), "NO_ACTIVE_STEPS"),
    (workbook(locator=":not(*)"), "UNRESOLVED_LOCATOR"), (b"not a workbook", "INVALID_WORKBOOK"),
    (workbook(action="invented"), "INVALID_ACTION")])
def test_gate_blocks_before_import_browser_or_environment(tmp_path, monkeypatch, content, code):
    source = tmp_path / "synthetic.xlsx"
    source.write_bytes(content)
    def forbidden(*args):
        raise AssertionError("Executor module must not be imported")
    monkeypatch.setattr(executor.importlib.util, "spec_from_file_location", forbidden)
    result = executor.execute(source, "run1", tmp_path / "out", env_path="must-not-be-opened")
    assert result["status"] == "ERROR" and result["errors"] == 1
    assert code in {issue["code"] for issue in result["preflight"]["issues"]}
    assert source.read_bytes() == content
    assert "synthetic-private-description" not in json.dumps(result)
    assert "example.test" not in json.dumps(result)
    assert json.loads((tmp_path / "out" / "preflight.json").read_text())["status"] == "blocked"


def test_ready_workbook_calls_executor_and_retains_static_warning(tmp_path, monkeypatch):
    source = tmp_path / "synthetic.xlsx"
    source.write_bytes(workbook())
    calls = []
    def run(config, **kwargs):
        calls.append((config, kwargs))
        summary = {"passed": 0, "failed": 0, "errors": 0, "unverified": 1, "duration": 0}
        (Path(kwargs["output_dir"]) / "summary.json").write_text(json.dumps(summary))
        return summary
    class Loader:
        def exec_module(self, module):
            module.execute_config = run
    monkeypatch.setattr(executor.importlib.util, "spec_from_file_location", lambda *_: SimpleNamespace(loader=Loader()))
    monkeypatch.setattr(executor.importlib.util, "module_from_spec", lambda *_: SimpleNamespace())
    executor.execute(source, "run1", tmp_path / "out")
    assert len(calls) == 1
    summary = json.loads((tmp_path / "out" / "summary.json").read_text())
    assert summary["preflight"]["status"] == "passed" and summary["unverified"] == 1
    assert summary["preflight"]["warnings"][0]["code"] == "NO_ACTIVE_ASSERTIONS"


def test_report_contract_rejects_values_unknown_codes_and_inconsistent_status():
    report = metadata(check(workbook(active="N")))
    for change in ({"message": "synthetic-private"}, {"status": "passed"},
                   {"issues": [{"sheet": "steps", "code": "arbitrary-text"}]}):
        with pytest.raises(ValueError):
            PreflightReport.model_validate({**report, **change})
    raw = check(workbook(active="N"))
    raw["issues"] = raw["issues"] * 101
    clipped = metadata(raw)
    assert len(clipped["issues"]) == 100 and clipped["truncated"] is True


def test_agent_reports_metadata_and_journal_resends_without_execution(tmp_path):
    reports = []
    def handler(req):
        reports.append(json.loads(req.content))
        return httpx.Response(503 if len(reports) == 1 else 200, json={})
    configs = tmp_path / "configs"
    configs.mkdir()
    (configs / "synthetic.xlsx").write_bytes(workbook())
    agent = LocalAgent("http://localhost", "synthetic", configs, tmp_path / "state",
                       client=httpx.Client(base_url="http://localhost", transport=httpx.MockTransport(handler)))
    report = metadata(check(workbook(active="N")))
    calls = []
    def worker(*_):
        calls.append(True)
        return {"passed": 0, "failed": 0, "errors": 1, "unverified": 0, "duration": 0, "preflight": report}
    run = {"run_id": "run1", "local_ref": "synthetic.xlsx"}
    agent.run(run, worker=worker)
    agent.run(run, worker=worker)
    agent.resend_results()
    assert len(calls) == 1 and len(reports) == 2 and reports[1]["preflight"] == report
    assert "synthetic-private-description" not in json.dumps(reports)


def test_preflight_action_and_locator_sets_match_current_runner():
    import ast
    tree = ast.parse((Path(__file__).resolve().parents[2] / "docs" / "runner.py").read_text(encoding="utf-8"))
    constants = {node.targets[0].id: ast.literal_eval(node.value)
        for node in tree.body if isinstance(node, ast.Assign) and isinstance(node.targets[0], ast.Name)
        and node.targets[0].id in {"VALID_ACTIONS", "VALID_LOCATOR_TYPES"}}
    assert VALID_ACTIONS == constants["VALID_ACTIONS"]
    assert VALID_LOCATOR_TYPES == constants["VALID_LOCATOR_TYPES"]
