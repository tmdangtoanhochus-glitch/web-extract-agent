import io
import json
from concurrent.futures import ThreadPoolExecutor

import httpx
import pytest
from fastapi import FastAPI
from fastapi.responses import JSONResponse
from fastapi.testclient import TestClient
from openpyxl import Workbook

from src.api.runner import create_runner_router
from src.runner.repository import Repository
from src.runner.service import RunnerError, Service, safe_path
from runner_agent.client import LocalAgent
from runner_agent.config import validate_workbook


@pytest.fixture
def system(tmp_path):
    repo = Repository()
    now = [1000000.0]
    service = Service(repo, tmp_path / "server", clock=lambda: now[0])
    owner = service.add_user("owner", "synthetic-password-1", "admin")
    other = service.add_user("other", "synthetic-password-2")
    pair = service.register_agent(owner)
    agent = service.agent(pair["agent_token"])
    app = FastAPI()
    @app.exception_handler(RunnerError)
    async def handler(req, exc):
        return JSONResponse(status_code=exc.code, content={"detail": str(exc)})
    app.include_router(create_runner_router(service))
    yield service, repo, now, owner, other, agent, TestClient(app)
    repo.close()


def login(client, name="owner", password="synthetic-password-1"):
    response = client.post("/runner/login", json={"username": name, "password": password})
    assert response.status_code == 200
    return {"Authorization": "Bearer " + response.json()["session"]}


def create(service, owner, agent):
    return service.create_run(owner, agent["id"], "UAT.xlsx", local_ref="UAT.xlsx")


def metrics():
    return {"passed": 1, "failed": 0, "errors": 0, "unverified": 0, "duration": 2}


def test_lost_run_retains_late_result_once_without_replay_or_changing_retention(system):
    s, repo, now, owner, _, agent, _ = system
    run = create(s, owner, agent)
    s.claim(agent)
    now[0] += 301
    s.maintenance()
    lost = repo.get("runs", run["run_id"])
    now[0] += 50
    result = s.finish(agent, run["run_id"], metrics())
    assert result["status"] == "LOST" and result["late_result"]["status"] == "PASSED"
    assert result["late_result"]["received_at"] == now[0]
    for field in ("finished_at", "expires_at", "notification_sent_at", "deleted_at"):
        assert result[field] == lost[field]
    first = result["late_result"]
    now[0] += 60
    assert s.finish(agent, run["run_id"], {**metrics(), "errors": 10})["late_result"] == first
    assert s.claim(agent) is None
    assert len([e for e in repo.all("audit") if e["event"] == "RUN_LATE_RESULT_RECEIVED_NO_REPLAY"]) == 1


def test_late_preflight_still_blocks_false_pass_and_never_reconciles_unstarted_run(system):
    s, repo, now, owner, _, agent, _ = system
    run = create(s, owner, agent)
    s.claim(agent)
    never_started = create(s, owner, agent)
    now[0] += 86401
    s.maintenance()
    report = {"status": "blocked", "active_steps": 0, "active_testcases": 0,
        "issues": [{"sheet": "steps", "code": "NO_ACTIVE_STEPS"}], "warnings": []}
    late = s.finish(agent, run["run_id"], {**metrics(), "preflight": report})["late_result"]
    assert late["status"] == "ERROR" and late["passed"] == 0 and late["errors"] == 1
    assert late["preflight"]["status"] == "blocked"
    assert "late_result" not in s.finish(agent, never_started["run_id"], metrics())


@pytest.mark.parametrize("expired", [False, True])
def test_late_result_api_owner_scope_summary_and_no_resurrection(system, expired):
    s, repo, now, owner, other, _, client = system
    pair = s.register_agent(owner)
    agent = s.agent(pair["agent_token"])
    run = create(s, owner, agent)
    s.claim(agent)
    now[0] += 301
    s.maintenance()
    if expired:
        now[0] += 6 * 86400
        s.maintenance()
        now[0] += 86400
        s.maintenance()
    headers = {"Authorization": "Bearer " + pair["agent_token"]}
    result_path = f"/runner/agent/runs/{run['run_id']}/result"
    stranger = s.register_agent(other)
    assert client.post(result_path, json=metrics(), headers={
        "Authorization": "Bearer " + stranger["agent_token"]}).status_code == 404
    response = client.post(result_path, json=metrics(), headers=headers)
    assert response.status_code == 200
    assert response.json()["late_result"]["passed"] == 1
    path = f"/runner/runs/{run['run_id']}"
    assert client.get(path, headers=login(client, "other", "synthetic-password-2")).status_code == 404
    summary = client.get(path + "/artifacts/summary.json", headers=login(client))
    if expired:
        assert summary.status_code == 404
        assert not (s.root / "runs" / run["run_id"]).exists()
        assert response.json()["deleted_at"] is not None
    else:
        assert summary.status_code == 200 and summary.json()["late_result"]["status"] == "PASSED"


def test_preflight_metadata_round_trip_is_owner_scoped_and_cannot_claim_pass(system):
    s, repo, _, owner, other, _, client = system
    pair = s.register_agent(owner)
    agent = s.agent(pair["agent_token"])
    run = create(s, owner, agent)
    s.claim(agent)
    report = {"status": "blocked", "active_steps": 0, "active_testcases": 0,
              "issues": [{"sheet": "testcases", "code": "ENTER_AND_ACTIVATE_USER_TESTCASES"}], "warnings": []}
    headers = {"Authorization": "Bearer " + pair["agent_token"]}
    path = f"/runner/agent/runs/{run['run_id']}/result"
    bad = client.post(path, json={**metrics(), "preflight": {**report, "message": "synthetic-private"}}, headers=headers)
    assert bad.status_code == 422
    response = client.post(path, json={**metrics(), "preflight": report}, headers=headers)
    assert response.status_code == 200 and response.json()["status"] == "ERROR" and response.json()["passed"] == 0
    own = login(client)
    saved = client.get(f"/runner/runs/{run['run_id']}", headers=own).json()
    assert saved["preflight"]["issues"][0]["code"] == "ENTER_AND_ACTIVATE_USER_TESTCASES"
    summary = client.get(f"/runner/runs/{run['run_id']}/artifacts/summary.json", headers=own).json()
    assert summary["preflight"]["status"] == "blocked"
    other_auth = login(client, "other", "synthetic-password-2")
    assert client.get(f"/runner/runs/{run['run_id']}", headers=other_auth).status_code == 404


def test_login_logout_and_hash_storage(system):
    s, repo, _, owner, _, _, client = system
    auth = login(client)
    assert client.get("/runner/me", headers=auth).status_code == 200
    with repo.transaction():
        raw = repo.get("users", owner["id"])
        assert raw["password_hash"].startswith("$argon2id$")
        assert "synthetic-password-1" not in json.dumps(raw)
    assert client.post("/runner/logout", headers=auth).status_code == 200
    assert client.get("/runner/me", headers=auth).status_code == 401


def test_wrong_password_and_unknown_user_rejected(system):
    client = system[-1]
    for username in ("owner", "missing"):
        assert client.post("/runner/login", json={"username": username, "password": "bad"}).status_code == 401


def test_owner_cannot_read_other_run_or_artifact(system):
    s, _, _, owner, _, agent, client = system
    run = create(s, owner, agent)
    auth = login(client, "other", "synthetic-password-2")
    assert client.get("/runner/runs", headers=auth).json() == []
    for suffix in ("", "/artifacts/summary.json"):
        assert client.get("/runner/runs/" + run["run_id"] + suffix, headers=auth).status_code == 404
    assert client.get("/runner/users", headers=auth).status_code == 403


def test_disabled_user_invalidates_existing_session_and_agent(system):
    s, _, _, _, other, _, client = system
    pair = s.register_agent(other)
    auth = login(client, "other", "synthetic-password-2")
    client.patch("/runner/users/" + other["id"], json={"is_active": False}, headers=login(client))
    assert client.get("/runner/me", headers=auth).status_code == 401
    assert client.post("/runner/agent/claim", headers={"Authorization": "Bearer " + pair["agent_token"]}).status_code == 401


def test_admin_can_reset_other_user_password(system):
    s, _, _, owner, other, _, client = system
    admin_auth = login(client)

    resp = client.post(
        "/runner/users/" + other["id"] + "/reset-password",
        json={"new_password": "brand-new-password-1"},
        headers=admin_auth,
    )
    assert resp.status_code == 200

    # Mật khẩu cũ không còn đăng nhập được, mật khẩu mới đăng nhập được.
    assert client.post("/runner/login", json={"username": "other", "password": "synthetic-password-2"}).status_code == 401
    assert client.post("/runner/login", json={"username": "other", "password": "brand-new-password-1"}).status_code == 200


def test_non_admin_cannot_reset_password(system):
    _, _, _, _, other, _, client = system
    auth = login(client, "other", "synthetic-password-2")

    resp = client.post(
        "/runner/users/" + other["id"] + "/reset-password",
        json={"new_password": "brand-new-password-1"},
        headers=auth,
    )
    assert resp.status_code == 403


def test_reset_password_rejects_short_password(system):
    _, _, _, _, other, _, client = system
    admin_auth = login(client)

    resp = client.post(
        "/runner/users/" + other["id"] + "/reset-password",
        json={"new_password": "short"},
        headers=admin_auth,
    )
    assert resp.status_code == 400


def test_forgot_password_creates_pending_request_for_admin(system):
    s, _, _, owner, other, _, client = system

    resp = client.post("/runner/forgot-password", json={"username": "other"})
    assert resp.status_code == 200

    admin_auth = login(client)
    pending = client.get("/runner/password-reset-requests", headers=admin_auth).json()
    assert len(pending) == 1
    assert pending[0]["username"] == "other"
    assert pending[0]["user_id"] == other["id"]


def test_forgot_password_does_not_leak_whether_username_exists(system):
    s, _, _, owner, other, _, client = system

    real_resp = client.post("/runner/forgot-password", json={"username": "other"})
    fake_resp = client.post("/runner/forgot-password", json={"username": "does-not-exist"})

    assert real_resp.status_code == fake_resp.status_code == 200
    assert real_resp.json() == fake_resp.json()

    admin_auth = login(client)
    pending = client.get("/runner/password-reset-requests", headers=admin_auth).json()
    assert len(pending) == 1  # chỉ request cho user thật tồn tại được tạo


def test_forgot_password_is_idempotent_for_repeated_requests(system):
    s, _, _, owner, other, _, client = system

    client.post("/runner/forgot-password", json={"username": "other"})
    client.post("/runner/forgot-password", json={"username": "other"})

    admin_auth = login(client)
    pending = client.get("/runner/password-reset-requests", headers=admin_auth).json()
    assert len(pending) == 1  # gửi 2 lần vẫn chỉ 1 request đang chờ


def test_reset_password_clears_pending_request_for_that_user(system):
    s, _, _, owner, other, _, client = system
    client.post("/runner/forgot-password", json={"username": "other"})
    admin_auth = login(client)

    client.post(
        "/runner/users/" + other["id"] + "/reset-password",
        json={"new_password": "brand-new-password-1"},
        headers=admin_auth,
    )

    pending = client.get("/runner/password-reset-requests", headers=admin_auth).json()
    assert pending == []


def test_password_reset_requests_route_requires_admin(system):
    _, _, _, _, other, _, client = system
    auth = login(client, "other", "synthetic-password-2")

    assert client.get("/runner/password-reset-requests", headers=auth).status_code == 403


def test_claim_is_atomic_and_never_reexecutes_active_job(system):
    s, _, _, owner, _, agent, _ = system
    run = create(s, owner, agent)
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(lambda _: s.claim(agent), range(4)))
    assert sum(r is not None for r in results) == 1
    assert next(r for r in results if r)["run_id"] == run["run_id"]


def test_run_name_collision_and_temp_deletion(system):
    s, _, _, owner, _, agent, _ = system
    first = s.create_run(owner, agent["id"], "UAT.xlsx", config=b"fake-only-service-test")
    second = create(s, owner, agent)
    assert first["run_id"] != second["run_id"]
    s.claim(agent)
    s.finish(agent, first["run_id"], metrics())
    assert not (s.root / "temp_uploads" / first["run_id"]).exists()


def test_result_is_idempotent_and_retention_preserves_audit(system):
    s, repo, now, owner, _, agent, _ = system
    run = create(s, owner, agent)
    s.claim(agent)
    done = s.finish(agent, run["run_id"], metrics())
    now[0] += 100
    assert s.finish(agent, run["run_id"], metrics())["finished_at"] == done["finished_at"]
    folder = s.root / "runs" / run["run_id"]
    folder.mkdir()
    (folder / "summary.json").write_text("{}")
    now[0] = done["expires_at"] - 86400
    s.maintenance()
    assert folder.exists()
    now[0] = done["expires_at"]
    s.maintenance()
    assert not folder.exists()
    with repo.transaction():
        r = repo.get("runs", run["run_id"])
        assert r["deleted_at"] == now[0]
        assert any(e["event"] == "RUN_ARTIFACTS_DELETED" for e in repo.all("audit"))


def test_late_notice_delays_deletion_and_offline_run_never_requeues(system):
    s, repo, now, owner, _, agent, _ = system
    r = create(s, owner, agent)
    s.claim(agent)
    now[0] += 301
    s.maintenance()
    with repo.transaction():
        assert repo.get("runs", r["run_id"])["status"] == "LOST"
    assert s.claim(agent) is None
    now[0] += 8 * 86400
    s.maintenance()
    with repo.transaction():
        assert repo.get("runs", r["run_id"])["deleted_at"] is None
    now[0] += 86400
    s.maintenance()
    with repo.transaction():
        assert repo.get("runs", r["run_id"])["deleted_at"] is not None


@pytest.mark.parametrize("name", ["../data", "a/b", "C:\\x", "..", "a:b"])
def test_paths_stay_inside_artifact_root(tmp_path, name):
    with pytest.raises(RunnerError):
        safe_path(tmp_path, name)


def workbook():
    wb = Workbook()
    wb.active.title = "settings"
    wb.active.append(["key", "value"])
    wb.active.append(["url", "https://example.test"])
    ws = wb.create_sheet("steps")
    ws.append(["screen", "step", "action", "value_source"])
    ws.append(["login", "username", "fill", "account"])
    ws = wb.create_sheet("testcases")
    ws.append(["tc_id", "role_code"])
    ws.append(["T1", "RM"])
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def test_agent_local_config_never_uploaded_and_never_deleted(tmp_path):
    configs = tmp_path / "configs"
    configs.mkdir()
    source = configs / "UAT.xlsx"
    source.write_bytes(workbook())
    seen = []
    def handler(req):
        seen.append((req.method, req.url.path, req.content))
        return httpx.Response(200, json={})
    client = httpx.Client(base_url="http://localhost:8000", transport=httpx.MockTransport(handler))
    agent = LocalAgent("http://localhost:8000", "fake", configs, tmp_path / "local", client=client)
    calls = []
    def worker(config, rid, output):
        calls.append(rid)
        assert config.read_bytes() == source.read_bytes()
        return metrics()
    run = {"run_id": "UAT_20260101_000000", "local_ref": "UAT.xlsx"}
    agent.run(run, worker=worker)
    agent.run(run, worker=worker)
    assert calls == [run["run_id"]]
    assert source.exists()
    assert not (agent.state / "temp_uploads" / run["run_id"]).exists()
    assert len(seen) == 1
    assert set(json.loads(seen[0][2])) == set(metrics())


def test_agent_blocks_remote_plain_http(tmp_path):
    with pytest.raises(ValueError):
        LocalAgent("http://remote.example", "fake", tmp_path, tmp_path)


def test_workbook_rejects_formula_and_plaintext_credential():
    for column, value in [("password", "synthetic-secret"), ("foo", "=1+1")]:
        wb = Workbook()
        wb.active.title = "settings"
        wb.create_sheet("steps")
        sheet = wb.create_sheet("testcases")
        sheet.append([column])
        sheet.append([value])
        stream = io.BytesIO()
        wb.save(stream)
        with pytest.raises((ValueError, StopIteration)):
            validate_workbook(stream.getvalue())
