import io
import json

import httpx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from runner_agent.authoring import _workbook
from src.api.runner import create_runner_router
from src.runner.planner import StepPlanner, PlanError
from src.runner.recording_plan import RecordingTrace, RecordingPlan
from src.runner.repository import Repository
from src.runner.service import Service


def evidence():
    return RecordingTrace(events=[
        {"id": 1, "screen": "recorded", "action": "fill", "locator": "input:nth-of-type(1)"},
        {"id": 2, "screen": "recorded", "action": "fill", "locator": "input:nth-of-type(1)"},
        {"id": 3, "screen": "recorded", "action": "click", "locator": "input:nth-of-type(2)", "widget": "antd_select"},
        {"id": 4, "screen": "recorded", "action": "click", "locator": "div:nth-of-type(3)",
         "widget": "antd_option", "related_locator": "input:nth-of-type(2)"},
        {"id": 5, "screen": "recorded", "action": "click", "locator": "button:nth-of-type(1)"},
        {"id": 6, "screen": "recorded", "action": "read_result_single", "locator": "input:nth-of-type(3)"},
    ])


def proposed():
    return {"steps": [
        {"event_ids": [1, 2], "screen": "recorded", "step": "record_code", "action": "fill", "value_source": "testcase"},
        {"event_ids": [3, 4], "screen": "recorded", "step": "unit", "action": "select_antd", "value_source": "testcase",
         "dropdown_selector": "visible", "match_type": "exact"},
        {"event_ids": [5], "screen": "recorded", "step": "save", "action": "click", "value_source": "empty"},
        {"event_ids": [6], "screen": "recorded", "step": "read_status", "action": "read_result_single", "value_source": "empty",
         "read_method": "css_input", "match_type": "exact"},
    ]}


def test_compiles_semantic_runner_actions_and_preserves_event_coverage():
    rows, notes = RecordingPlan.model_validate(proposed()).bind(evidence())
    assert len(rows) == 4
    assert rows[1]["action"] == "select_antd" and rows[1]["locator"] == "input:nth-of-type(2)"
    assert rows[1]["match_type"] == "exact" and rows[1]["dropdown_selector"] == "visible"
    wb = load_workbook(io.BytesIO(_workbook(rows, compilation_notes=notes)))
    try:
        assert list(next(wb["testcases"].values))[-3:] == ["record_code", "unit", "expected_status"]
        assert wb["testcases"].max_row == 1 and "settings" not in wb.sheetnames
        assert all(row[6] == "N" for row in list(wb["steps"].values)[1:])
        assert any("events=3,4" in str(row) for row in wb["review"].values)
    finally:
        wb.close()


@pytest.mark.parametrize("change", ["omit", "reorder", "duplicate", "account", "wrong_action", "invented_locator", "activate", "group", "screen", "wait", "wrong_dropdown"])
def test_rejects_unproven_changes(change):
    data = proposed()
    steps = data["steps"]
    if change == "omit": steps.pop(2)
    if change == "reorder": steps[0], steps[1] = steps[1], steps[0]
    if change == "duplicate": steps[0]["event_ids"] = [1, 1, 2]
    if change == "account": steps[0].update(step="username", value_source="account")
    if change == "wrong_action": steps[2]["action"] = "check"
    if change == "invented_locator": steps[0]["locator"] = "input:nth-of-type(9)"
    if change == "activate": steps[0]["active"] = "Y"
    if change == "group": steps[0]["group"] = "guessed"
    if change == "screen": steps[0]["screen"] = "guessed"
    if change == "wait": steps[0]["wait_selector"] = ":not(*)"
    if change == "wrong_dropdown": steps[1]["dropdown_selector"] = "not_hidden"
    with pytest.raises(ValueError):
        RecordingPlan.model_validate(data).bind(evidence())


def test_missing_component_evidence_cannot_be_upgraded_to_antd():
    trace = evidence().model_dump()
    trace["events"][2]["widget"] = "native"
    with pytest.raises(ValueError):
        RecordingPlan.model_validate(proposed()).bind(RecordingTrace.model_validate(trace))


def test_input_button_clicks_survive_compilation_without_becoming_testcase_inputs():
    from runner_agent.authoring import Recording
    recording = Recording()
    for _ in range(2):
        assert recording.accept({"action": "click", "locator": "input:nth-of-type(2)"})
    trace = RecordingTrace(events=[{"id": i, **event} for i, event in enumerate(recording.events, 1)])
    plan = {"steps": [
        {"event_ids": [i], "screen": "recorded", "step": f"press_{i}",
         "action": "click", "value_source": "empty"} for i in (1, 2)
    ]}
    rows, notes = RecordingPlan.model_validate(plan).bind(trace)
    wb = load_workbook(io.BytesIO(_workbook(rows, compilation_notes=notes)))
    try:
        assert [row[2] for row in list(wb["steps"].values)[1:]] == ["click", "click"]
        assert all(row[4:7] == ("input:nth-of-type(2)", "empty", "N")
                   for row in list(wb["steps"].values)[1:])
        assert not {"press_1", "press_2"} & set(next(wb["testcases"].values))
        assert wb["testcases"].max_row == 1 and "settings" not in wb.sheetnames
    finally:
        wb.close()
    # Two actual clicks must not be collapsed like consecutive fills.
    plan["steps"] = [{**plan["steps"][0], "event_ids": [1, 2]}]
    with pytest.raises(ValueError, match="Only consecutive fills"):
        RecordingPlan.model_validate(plan).bind(trace)


def test_uncertain_locator_remains_placeholder_and_repeat_is_review_only():
    data = proposed()
    data["steps"][0]["review"] = "REVIEW_LOCATOR"
    data["steps"][1]["review"] = "REVIEW_REPEAT"
    rows, notes = RecordingPlan.model_validate(data).bind(evidence())
    assert rows[0]["locator"] == ":not(*)"
    assert rows[1]["group"] == "" and notes[1]["review"] == "REVIEW_REPEAT"


@pytest.mark.parametrize("field", ["value", "filename", "text", "url", "cookie"])
def test_trace_rejects_page_data(field):
    data = evidence().model_dump()
    data["events"][0][field] = "synthetic-private"
    with pytest.raises(ValueError):
        RecordingTrace.model_validate(data)


def test_planner_sends_only_reviewed_evidence_and_rejects_invalid_mapping():
    calls = []
    def handler(request):
        calls.append(json.loads(request.content))
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(proposed())}}]})
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        planner = StepPlanner("https://ai.example/v1", "synthetic", "synthetic-model", client=client)
        plan = planner.recording("Nhập mã hồ sơ rồi chọn đơn vị và lưu, đọc trạng thái.", evidence())
        assert len(plan.steps) == 4
        assert json.loads(calls[0]["messages"][1]["content"])["recording"] == evidence().model_dump()
        with pytest.raises(PlanError):
            planner.recording("password=synthetic-private", evidence())
        assert len(calls) == 1


def test_api_requires_login_review_and_never_persists_trace(tmp_path):
    calls = []
    class Planner:
        def recording(self, description, trace):
            calls.append(trace)
            return RecordingPlan.model_validate(proposed())
    repo = Repository()
    try:
        service = Service(repo, tmp_path / "server")
        service.add_user("tester", "synthetic-test-password")
        session = service.login("tester", "synthetic-test-password")["session"]
        app = FastAPI()
        app.include_router(create_runner_router(service, Planner()))
        with TestClient(app) as client:
            body = {"description": "Nhập mã rồi chọn đơn vị và lưu.", "recording": evidence().model_dump(), "reviewed_no_secrets": True}
            assert client.post("/runner/authoring/recording", json=body).status_code == 401
            headers = {"Authorization": "Bearer " + session}
            assert client.post("/runner/authoring/recording", headers=headers, json={**body, "reviewed_no_secrets": False}).status_code == 422
            result = client.post("/runner/authoring/recording", headers=headers, json=body)
            assert result.status_code == 200 and result.headers["cache-control"] == "no-store"
            assert len(calls) == 1 and not repo.all("runs")
            audits = repo.all("audit")
            assert any(a["event"] == "RECORDING_DRAFT_COMPILED" for a in audits)
            assert "nth-of-type" not in json.dumps(audits)
            assert not list((tmp_path / "server").rglob("*.json"))
    finally:
        repo.close()


def test_recorder_exports_structural_trace_for_ai_without_changing_raw_draft(tmp_path, monkeypatch):
    from types import SimpleNamespace
    import playwright.sync_api
    from record_runner import record
    class Context:
        pages = []
        def expose_binding(self, name, callback): self.callback = callback
        def add_init_script(self, **kwargs): pass
        def new_page(self):
            page = SimpleNamespace(main_frame=object())
            for event in evidence().events:
                payload = event.model_dump(exclude={"id", "screen"}, exclude_defaults=True)
                self.callback({"page": page, "frame": page.main_frame}, payload)
            return page
    browser = SimpleNamespace(new_context=Context, is_connected=lambda: True, close=lambda: None)
    class Playwright:
        def __enter__(self): return SimpleNamespace(chromium=SimpleNamespace(launch=lambda **kw: browser))
        def __exit__(self, *args): pass
    monkeypatch.setattr(playwright.sync_api, "sync_playwright", Playwright)
    workbook, trace_path = tmp_path / "raw.xlsx", tmp_path / "events.json"
    record(workbook, trace_path)
    trace = RecordingTrace.model_validate_json(trace_path.read_text())
    expected = evidence().model_dump()
    expected["events"].pop(1)  # Consecutive typing is now coalesced at the shared receiver.
    for i, event in enumerate(expected["events"], 1):
        event["id"] = i
    assert trace == RecordingTrace.model_validate(expected)
    wb = load_workbook(workbook)
    assert wb["steps"].max_row == 6 and wb["testcases"].max_row == 1
    assert [row[2] for row in list(wb["steps"].values)[1:]] == [event.action for event in trace.events]
    wb.close()


def test_compose_keeps_event_mapping_and_review_instructions(tmp_path):
    from compose_runner import compose
    rows, notes = RecordingPlan.model_validate(proposed()).bind(evidence())
    source, output = tmp_path / "source.xlsx", tmp_path / "composed.xlsx"
    source.write_bytes(_workbook(rows, compilation_notes=notes))
    compose([source], output)
    wb = load_workbook(output)
    try:
        assert any("events=3,4" in str(row) for row in wb["review"].values)
        assert wb["testcases"].max_row == 1
    finally:
        wb.close()
