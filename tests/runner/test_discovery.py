import io
import json
from types import SimpleNamespace

import httpx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api.runner import create_runner_router
from src.runner.discovery import Snapshot, DiscoveryPlan, RepairChoice, repair_proposal
from src.runner.planner import StepPlanner, PlanError
from src.runner.repository import Repository
from src.runner.service import Service
from runner_agent.authoring import Recording, draft_workbook, _workbook
from runner_agent.discovery import selected_candidate, apply_proposal

RAW = {"candidates": [
    {"id": "c1", "selector": "html:nth-of-type(1) > body:nth-of-type(1) > input:nth-of-type(1)", "kind": "input"},
    {"id": "c2", "selector": "html:nth-of-type(1) > body:nth-of-type(1) > button:nth-of-type(1)", "kind": "button"}]}
PLAN = {"steps": [
    {"screen": "login", "step": "username", "action": "fill", "value_source": "account", "candidate_id": "c1"},
    {"screen": "login", "step": "submit", "action": "click", "value_source": "empty", "candidate_id": "c2"}]}
DESCRIPTION = "Điền tài khoản từ account vào c1 rồi click c2."


def planner(payload, requests=None):
    def handle(request):
        if requests is not None:
            requests.append(json.loads(request.content))
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(payload)}}]})
    return StepPlanner("https://ai.example/v1", "synthetic-key", "synthetic-model",
                       client=httpx.Client(transport=httpx.MockTransport(handle)))


@pytest.mark.parametrize("change", [
    {"text": "synthetic-private"}, {"value": "synthetic-private"},
    {"selector": "input[value='synthetic-private']"},
    {"selector": "private-custom-tag:nth-of-type(1)"},
    {"kind": "button"}, {"id": "synthetic-private"},
])
def test_snapshot_rejects_values_attributes_custom_tags_and_incoherent_metadata(change):
    data = {"candidates": [{**RAW["candidates"][0], **change}]}
    with pytest.raises(ValueError):
        Snapshot.model_validate(data)


def test_snapshot_unique_size_and_unknown_fields():
    for data in ({**RAW, "url": "https://synthetic.example"}, {"candidates": RAW["candidates"] * 2},
                 {"candidates": []}, {"candidates": RAW["candidates"] * 51}):
        with pytest.raises(ValueError):
            Snapshot.model_validate(data)


def test_discovery_uses_only_reviewed_candidates_and_exports_no_testcases_or_settings():
    snapshot = Snapshot.model_validate(RAW)
    requests = []
    plan = planner(PLAN, requests).discover(DESCRIPTION, snapshot)
    wb = load_workbook(io.BytesIO(_workbook(plan.bind(snapshot))))
    assert wb["testcases"].max_row == 1
    assert "settings" not in wb.sheetnames
    rows = list(wb["steps"].values)
    assert len(rows[0]) == 13
    assert rows[1][4] == RAW["candidates"][0]["selector"]
    assert rows[1][6] == rows[2][6] == "N"
    wb.close()
    sent = json.loads(requests[0]["messages"][1]["content"])
    assert sent == {"description": DESCRIPTION, "snapshot": snapshot.model_dump()}


@pytest.mark.parametrize("change", [
    {"candidate_id": "c9"}, {"candidate_id": "c2"}, {"locator": "input"},
    {"active": "Y"}, {"value": "synthetic-private"},
])
def test_model_cannot_invent_target_activate_or_supply_data(change):
    payload = {"steps": [{**PLAN["steps"][0], **change}]}
    with pytest.raises(PlanError):
        planner(payload).discover(DESCRIPTION, Snapshot.model_validate(RAW))


def test_unmentioned_candidates_and_multi_screen_are_rejected():
    with pytest.raises(PlanError):
        planner(PLAN).discover("Điền dữ liệu và đăng nhập.", Snapshot.model_validate(RAW))
    plan = DiscoveryPlan.model_validate({"steps": [PLAN["steps"][0], {**PLAN["steps"][1], "screen": "next"}]})
    with pytest.raises(ValueError, match="one screen"):
        plan.bind(Snapshot.model_validate(RAW))


def test_api_auth_consent_feature_flag_audit_and_no_content_persistence(tmp_path):
    repo = Repository()
    service = Service(repo, tmp_path / "server")
    service.add_user("tester", "synthetic-password-123")
    auth = {"Authorization": "Bearer " + service.login("tester", "synthetic-password-123")["session"]}
    for mode, payload in [("discover", PLAN), ("repair", {"candidate_id": "c1", "reason": "USER_IDENTIFIED_TARGET"})]:
        body = {"description": DESCRIPTION, "snapshot": RAW, "reviewed_no_secrets": True}
        if mode == "repair":
            body["action"] = "fill"
        for enabled in (False, True):
            app = FastAPI()
            app.include_router(create_runner_router(service, planner(payload) if enabled else None))
            client = TestClient(app)
            path = "/runner/authoring/" + mode
            assert client.post(path, json=body).status_code == 401
            assert client.post(path, json={**body, "reviewed_no_secrets": False}, headers=auth).status_code == 422
            result = client.post(path, json=body, headers=auth)
            assert result.status_code == (200 if enabled else 503)
            if enabled:
                assert result.headers["cache-control"] == "no-store"
                if mode == "repair":
                    assert result.json()["snapshot_sha256"] == Snapshot.model_validate(RAW).fingerprint()
    with repo.transaction():
        assert repo.all("runs") == []
        audit = json.dumps(repo.all("audit"), ensure_ascii=False)
    assert "DISCOVERY_DRAFT_GENERATED" in audit and "REPAIR_PROPOSED" in audit
    assert DESCRIPTION not in audit and "selector" not in audit
    repo.close()


def test_repair_proposal_rejects_wrong_snapshot_action_and_unmentioned_id():
    snapshot = Snapshot.model_validate(RAW)
    choice = RepairChoice(candidate_id="c1", reason="USER_IDENTIFIED_TARGET")
    proposal = repair_proposal(choice, snapshot, "fill")
    assert selected_candidate(snapshot, proposal, {"action": "fill"}).id == "c1"
    with pytest.raises(ValueError):
        selected_candidate(snapshot, proposal, {"action": "click"})
    changed = Snapshot.model_validate({"candidates": RAW["candidates"][:1]})
    with pytest.raises(ValueError):
        selected_candidate(changed, proposal, {"action": "fill"})
    with pytest.raises(PlanError):
        planner(choice.model_dump()).repair("Chọn phần tử đăng nhập.", snapshot, "fill")


def fake_browser(monkeypatch, page):
    import playwright.sync_api
    closed = []
    context = SimpleNamespace(pages=[page], new_page=lambda: page)
    browser = SimpleNamespace(new_context=lambda: context, is_connected=lambda: True, close=lambda: closed.append(True))
    class Playwright:
        def __enter__(self): return SimpleNamespace(chromium=SimpleNamespace(launch=lambda **kw: browser))
        def __exit__(self, *args): pass
    monkeypatch.setattr(playwright.sync_api, "sync_playwright", Playwright)
    return context, closed


@pytest.mark.parametrize("finish", ["EXPORT", "q", "stale"])
def test_local_ai_repair_requires_live_confirmation_and_preserves_source(tmp_path, monkeypatch, finish):
    snapshot = Snapshot.model_validate(RAW)
    proposal = repair_proposal(RepairChoice(candidate_id="c1", reason="USER_IDENTIFIED_TARGET"), snapshot, "fill")
    snapshot_path, proposal_path = tmp_path / "snapshot.json", tmp_path / "proposal.json"
    snapshot_path.write_text(snapshot.model_dump_json())
    proposal_path.write_text(proposal.model_dump_json())
    recording = Recording()
    recording.accept({"action": "fill", "locator": "input:nth-of-type(2)"})
    source, output = tmp_path / "source.xlsx", tmp_path / "output.xlsx"
    original = draft_workbook(recording)
    source.write_bytes(original)
    live, highlighted = {"identity": True}, []
    locator = SimpleNamespace(count=lambda: 1, is_visible=lambda: True,
        evaluate=lambda script: live["identity"], highlight=lambda: highlighted.append(True))
    context, closed = fake_browser(monkeypatch, SimpleNamespace(locator=lambda s: locator))
    commands = iter(["1", "EXPORT" if finish == "stale" else finish])
    def answer(_):
        result = next(commands)
        if result == "EXPORT" and finish == "stale":
            live["identity"] = False
            context.pages = []
        return result
    result = apply_proposal(source, 2, output, snapshot_path, proposal_path, prompt=answer)
    assert bool(result) == output.exists() == (finish == "EXPORT")
    assert highlighted == [True] and closed == [True]
    assert source.read_bytes() == original
    if output.exists():
        wb = load_workbook(output)
        assert wb["steps"].cell(2, 5).value == RAW["candidates"][0]["selector"]
        assert wb["steps"].cell(2, 7).value == "N"
        assert wb["testcases"].max_row == 1
        wb.close()


@pytest.mark.parametrize("finish", ["EXPORT", "q", "stale"])
def test_discovery_cli_requires_review_and_refuses_changed_structure(tmp_path, monkeypatch, finish):
    from discover_runner import discover
    live = {"data": RAW}
    highlighted = []
    page = SimpleNamespace(evaluate=lambda script: live["data"],
        locator=lambda s: SimpleNamespace(highlight=lambda: highlighted.append(s)))
    _, closed = fake_browser(monkeypatch, page)
    commands = iter(["1", "c1", "EXPORT" if finish == "stale" else finish, "q"])
    def answer(_):
        command = next(commands)
        if command == "EXPORT" and finish == "stale":
            live["data"] = {"candidates": RAW["candidates"][:1]}
        return command
    output = tmp_path / "snapshot.json"
    result = discover(output, prompt=answer)
    assert bool(result) == output.exists() == (finish == "EXPORT")
    assert highlighted == [RAW["candidates"][0]["selector"]] and closed == [True]
    if output.exists():
        assert Snapshot.model_validate_json(output.read_bytes()) == Snapshot.model_validate(RAW)
        with pytest.raises(ValueError):
            discover(output)
