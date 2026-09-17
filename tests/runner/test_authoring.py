import io

from openpyxl import load_workbook
import pytest

from runner_agent.authoring import Recording, draft_workbook
from runner_agent.config import validate_workbook


@pytest.mark.parametrize("payload", [
    {"action": "fill", "locator": "input:nth-of-type(1)", "value": "synthetic-private"},
    {"action": "click", "locator": "[value='synthetic-private']"},
    {"action": "click", "locator": "=HYPERLINK(1)"},
    {"action": "upload", "locator": "input:nth-of-type(1)"},
    {"action": "fill", "locator": "input:nth-of-type(0)"},
    {"action": "click", "locator": 42},
    None,
])
def test_rejects_values_arbitrary_locators_and_unsupported_events(payload):
    recording = Recording()
    assert not recording.accept(payload)
    assert recording.events == []
    assert recording.dropped == 1


def test_draft_is_inactive_and_has_local_references_only():
    recording = Recording()
    for action in ("fill", "click", "select"):
        assert recording.accept({"action": action, "locator": "html:nth-of-type(1) > body:nth-of-type(1)"})
    content = draft_workbook(recording)
    validate_workbook(content, require_settings=False)
    with pytest.raises(ValueError, match="Missing sheets"):
        validate_workbook(content)
    wb = load_workbook(io.BytesIO(content))
    steps = list(wb["steps"].values)
    assert all(row[6] == "N" for row in steps[1:])
    assert wb["testcases"].max_row == 1
    assert list(wb["testcases"].values)[0][5:] == ("field_0001", "field_0003")
    assert "settings" not in wb.sheetnames
    wb.close()


def test_capture_limit_and_mutated_event_revalidation():
    recording = Recording(limit=1)
    assert recording.accept({"action": "click", "locator": "button:nth-of-type(1)"})
    assert not recording.accept({"action": "click", "locator": "button:nth-of-type(2)"})
    recording.events[0]["value"] = "synthetic-private"
    with pytest.raises(ValueError, match="Invalid recording"):
        draft_workbook(recording)


def test_record_controls_screens_pause_wait_and_expected_headers_only():
    recording = Recording()
    assert recording.accept({"action": "fill", "locator": "input:nth-of-type(1)"})
    assert recording.control({"control": "toggle_pause"})
    assert not recording.accept({"action": "click", "locator": "button:nth-of-type(1)"})
    assert recording.suppressed == 1 and recording.dropped == 0
    assert recording.control({"control": "next_screen"})
    assert recording.control({"control": "toggle_pause"})
    assert recording.accept({"action": "wait", "locator": "input:nth-of-type(1)"})
    assert recording.accept({"action": "read_result_single", "locator": "input:nth-of-type(1)"})
    wb = load_workbook(io.BytesIO(draft_workbook(recording)))
    rows = list(wb["steps"].values)
    assert [row[0] for row in rows[1:]] == ["recorded", "recorded_002", "recorded_002"]
    assert rows[2][7] == "input:nth-of-type(1)"
    assert rows[3][9] == "exact" and rows[3][12] == "css_input"
    assert all(row[6] == "N" for row in rows[1:])
    assert list(next(wb["testcases"].values))[-2:] == ["field_0001", "expected_field_0003"]
    assert wb["testcases"].max_row == 1 and "settings" not in wb.sheetnames
    wb.close()


def test_record_rejects_untrusted_controls_cross_screen_reads_and_wrong_read_targets():
    recording = Recording()
    for payload in ({"control": "navigate"}, {"control": "next_screen", "url": "synthetic-private"},
                    {"control": "next_screen", "screen": "private"}, None):
        assert not recording.control(payload)
    assert recording.screen == "recorded"
    assert not recording.accept({"action": "read_result_single", "locator": "button:nth-of-type(1)"})
    assert recording.accept({"action": "read_result_single", "locator": "textarea:nth-of-type(1)"})
    assert recording.control({"control": "next_screen"})
    assert not recording.accept({"action": "read_result_single", "locator": "input:nth-of-type(1)"})
    assert not recording.accept({"action": "fill", "locator": "input:nth-of-type(1)", "screen": "private"})
    for _ in range(98):
        assert recording.control({"control": "next_screen"})
    assert not recording.control({"control": "next_screen"})
    assert recording.screen == "recorded_100"


def test_invalid_export_does_not_leave_empty_output(tmp_path, monkeypatch):
    from types import SimpleNamespace
    import playwright.sync_api
    import record_runner
    class Playwright:
        def __enter__(self):
            context = SimpleNamespace(pages=[], expose_binding=lambda *a: None,
                add_init_script=lambda **kw: None, new_page=lambda: None)
            browser = SimpleNamespace(new_context=lambda: context, is_connected=lambda: True, close=lambda: None)
            return SimpleNamespace(chromium=SimpleNamespace(launch=lambda **kw: browser))
        def __exit__(self, *args): pass
    monkeypatch.setattr(playwright.sync_api, "sync_playwright", Playwright)
    def invalid(_): raise ValueError("Invalid draft")
    monkeypatch.setattr(record_runner, "draft_workbook", invalid)
    target = tmp_path / "out.xlsx"
    with pytest.raises(ValueError):
        record_runner.record(target)
    assert not target.exists()


def test_recorder_does_not_overwrite_existing_output(tmp_path):
    from record_runner import record
    path = tmp_path / "draft.xlsx"
    path.write_bytes(b"existing")
    with pytest.raises(ValueError, match="already exists"):
        record(path)
    assert path.read_bytes() == b"existing"


def test_record_session_exports_only_main_frame_events(tmp_path, monkeypatch):
    from types import SimpleNamespace
    import playwright.sync_api
    from record_runner import record

    closed = []
    class Context:
        pages = []
        def expose_binding(self, name, callback):
            self.callback = callback
        def add_init_script(self, script):
            assert "runnerRecordEvent" in script
        def new_page(self):
            page = SimpleNamespace(main_frame=object())
            self.callback({"page": page, "frame": object()},
                          {"action": "click", "locator": "button:nth-of-type(2)"})
            self.callback({"page": page, "frame": page.main_frame},
                          {"action": "fill", "locator": "input:nth-of-type(1)"})
            return page
    class Browser:
        def new_context(self): return Context()
        def is_connected(self): return True
        def close(self): closed.append(True)
    class Playwright:
        def __enter__(self):
            return SimpleNamespace(chromium=SimpleNamespace(launch=lambda **kw: Browser()))
        def __exit__(self, *args): pass
    monkeypatch.setattr(playwright.sync_api, "sync_playwright", Playwright)
    path = tmp_path / "draft.xlsx"
    record(path)
    wb = load_workbook(path)
    assert wb["steps"].max_row == 2
    assert wb["steps"].cell(2, 3).value == "fill"
    assert closed == [True]
    wb.close()
