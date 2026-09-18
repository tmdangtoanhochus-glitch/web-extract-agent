import io
import importlib.util
from pathlib import Path
from openpyxl import Workbook, load_workbook
import pytest
from pydantic import ValidationError
from src.runner.workbook_schema import STEP_COLUMNS, WorkbookPlan
from runner_agent.authoring import planned_workbook
from runner_agent.preparation import analyze, merge


def plan():
    return WorkbookPlan.model_validate({"steps": [
        {"screen": "login", "step": "username", "action": "fill", "value_source": "account"},
        {"screen": "search", "step": "cif", "action": "fill_enter", "value_source": "testcase"},
        {"screen": "result", "step": "read_amount", "action": "read_result_single",
         "value_source": "empty", "read_method": "css_input"}]})


def template(flow="login,search,result", result="result"):
    wb = Workbook()
    wb.active.title = "settings"
    for row in [("key", "value"), ("url", "https://example.test"), ("screen_flow", flow),
                ("result_screen", result), ("default_timeout", "4321"), ("popup_button", ".original-popup")]:
        wb.active.append(row)
    wb.create_sheet("steps").append(STEP_COLUMNS)
    cases = wb.create_sheet("testcases")
    cases.append(["tc_id", "mo_ta", "active", "role_code", "specialized_bank", "custom"])
    cases.append(["USER_01", "Human entered synthetic case", "N", "RM", "", "KEEP_ME"])
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def test_generation_only_creates_steps_and_testcase_headers():
    wb = load_workbook(io.BytesIO(planned_workbook(plan())))
    assert wb.sheetnames == ["steps", "testcases", "review"]
    assert list(next(wb["steps"].values)) == STEP_COLUMNS
    assert list(next(wb["testcases"].values)) == ["tc_id", "mo_ta", "active", "role_code", "specialized_bank", "cif", "expected_amount"]
    assert wb["testcases"].max_row == 1
    wb.close()


def test_prepare_preserves_compiler_review_without_overwriting_user_notes():
    from runner_agent.authoring import _workbook
    source = load_workbook(io.BytesIO(template()))
    source.create_sheet("draft_review").append(["original user note"])
    saved = io.BytesIO()
    source.save(saved)
    source.close()
    draft = _workbook([step.model_dump() for step in plan().steps], compilation_notes=[
        {"step": "cif", "events": [1, 2], "review": "REVIEW_LOCATOR"}])
    result = load_workbook(io.BytesIO(merge(saved.getvalue(), draft)))
    try:
        assert result["draft_review"].cell(1, 1).value == "original user note"
        assert any("events=1,2; REVIEW_LOCATOR" in str(row) for row in result["draft_review1"].values)
        assert result["testcases"].cell(2, 6).value == "KEEP_ME"
        assert all(row[6] == "N" for row in list(result["steps"].values)[1:])
    finally:
        result.close()


@pytest.mark.parametrize("key,value", [("settings", {}), ("testcases", []), ("testcases", [{"tc_id": "FAKE"}])])
def test_model_cannot_generate_settings_or_any_testcases(key, value):
    with pytest.raises(ValidationError):
        WorkbookPlan.model_validate({**plan().model_dump(), key: value})


@pytest.mark.parametrize("changes", [{"active": "Y"}, {"locator": "#guessed"}, {"step": "tc_id", "value_source": "testcase"}])
def test_unsafe_step_rejected(changes):
    data = plan().model_dump()
    data["steps"][0].update(changes)
    with pytest.raises(ValidationError):
        WorkbookPlan.model_validate(data)


def test_merge_preserves_settings_and_user_rows_and_adds_only_missing_headers():
    source = template()
    draft = planned_workbook(plan())
    _, proposals = analyze(source, draft)
    assert proposals == []
    before = load_workbook(io.BytesIO(source))
    after = load_workbook(io.BytesIO(merge(source, draft)))
    assert list(before["settings"].values) == list(after["settings"].values)
    assert list(after["testcases"].values)[1][:6] == list(before["testcases"].values)[1]
    assert after["testcases"].max_row == 2
    assert list(after["testcases"].values)[1][6:] == (None, None)
    assert all(row[6] == "N" for row in list(after["steps"].values)[1:])
    before.close()
    after.close()


def test_only_explained_and_approved_settings_can_change():
    source = template(flow="login", result="old_result")
    draft = planned_workbook(plan())
    _, proposals = analyze(source, draft)
    assert {p["key"] for p in proposals} == {"screen_flow", "result_screen"}
    assert all(p["reason"] for p in proposals)
    rejected = load_workbook(io.BytesIO(merge(source, draft)))
    assert dict(list(rejected["settings"].values)[1:])["screen_flow"] == "login"
    accepted = load_workbook(io.BytesIO(merge(source, draft, ["result_screen"])))
    settings = dict(list(accepted["settings"].values)[1:])
    assert settings["result_screen"] == "result" and settings["screen_flow"] == "login"
    assert settings["default_timeout"] == "4321" and settings["popup_button"] == ".original-popup"
    with pytest.raises(ValueError, match="explained"):
        merge(source, draft, ["url"])
    accepted.close()
    rejected.close()


def test_prepare_cli_defaults_no_and_never_changes_source(tmp_path):
    from prepare_runner import prepare
    source = tmp_path / "source.xlsx"
    draft = tmp_path / "draft.xlsx"
    target = tmp_path / "output.xlsx"
    content = template(flow="login", result="old_result")
    source.write_bytes(content)
    draft.write_bytes(planned_workbook(plan()))
    prompts = []
    prepare(source, draft, target, prompt=lambda text: (prompts.append(text) or ""))
    assert len(prompts) == 2
    assert source.read_bytes() == content
    wb = load_workbook(target)
    assert dict(list(wb["settings"].values)[1:])["result_screen"] == "old_result"
    wb.close()


def test_user_filled_config_loads_in_actual_runner(tmp_path):
    path = Path(__file__).resolve().parents[2] / "docs" / "runner.py"
    spec = importlib.util.spec_from_file_location("prepared_runner_test", path)
    runner = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(runner)
    wb = load_workbook(io.BytesIO(merge(template(), planned_workbook(plan()))))
    # Emulate the HUMAN filling cells after preparation; generator never writes these.
    wb["testcases"].cell(2, 3, "Y")
    wb["testcases"].cell(2, 7, "${FAKE_CIF}")
    wb["testcases"].cell(2, 8, "${FAKE_AMOUNT}")
    for row in range(2, wb["steps"].max_row + 1):
        wb["steps"].cell(row, 7, "Y")
    file = tmp_path / "prepared.xlsx"
    wb.save(file)
    wb.close()
    assert runner.validate_config(runner.load_steps(file), runner.load_testcases(file), runner.load_settings(file))


def test_preflight_reports_empty_user_testcases_without_copying_data():
    from runner_agent.preflight import check
    report = check(planned_workbook(plan()))
    codes = {item["code"] for item in report["issues"]}
    assert "ENTER_AND_ACTIVATE_USER_TESTCASES" in codes
    assert "MISSING_SETTINGS_USE_PREPARE" in codes
    assert not report["ready_for_local_review"]


def test_preflight_detects_placeholders_and_screen_settings():
    import json
    from runner_agent.preflight import check
    wb = load_workbook(io.BytesIO(merge(template(flow="login", result="old_result"), planned_workbook(plan()))))
    for row in range(2, wb["steps"].max_row + 1):
        wb["steps"].cell(row, 7, "Y")
    wb["testcases"].cell(2, 3, "Y")
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    report = check(out.getvalue())
    codes = {item["code"] for item in report["issues"]}
    assert {"UNRESOLVED_LOCATOR", "RESULT_SCREEN_MISMATCH", "SCREEN_NOT_IN_FLOW"} <= codes
    assert "KEEP_ME" not in json.dumps(report)
    assert "example.test" not in json.dumps(report)


def test_prepare_cli_applies_only_explicit_yes(tmp_path):
    from prepare_runner import prepare
    source, draft, output = (tmp_path / name for name in ("source.xlsx", "draft.xlsx", "out.xlsx"))
    source.write_bytes(template(flow="login", result="old_result"))
    draft.write_bytes(planned_workbook(plan()))
    answers = iter(["n", "y"])
    prepare(source, draft, output, prompt=lambda _: next(answers))
    wb = load_workbook(output)
    settings = dict(list(wb["settings"].values)[1:])
    assert settings["screen_flow"] == "login" and settings["result_screen"] == "result"
    assert wb["testcases"].max_row == 2
    wb.close()
