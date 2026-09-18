import importlib.util
import io
from pathlib import Path

from openpyxl import load_workbook
import pytest

from runner_agent.authoring import planned_workbook
from runner_agent.preparation import merge
from src.runner.workbook_schema import WorkbookPlan, header_result_blocks


def plan():
    return WorkbookPlan.model_validate({"steps": [
        {"screen": "entry", "step": "amount", "action": "fill", "value_source": "testcase", "group": "items"},
        {"screen": "entry", "step": "add", "action": "click", "value_source": "empty", "group": "items"},
        {"screen": "result", "step": "read_amount", "action": "read_result_group", "value_source": "empty", "read_method": "css_input"},
    ]})


def test_repeat_steps_and_multiple_expected_headers_without_testcase_or_settings(tmp_path):
    from compose_runner import compose
    generated = plan()
    content = planned_workbook(generated, result_blocks=3)
    wb = load_workbook(io.BytesIO(content))
    headers = list(next(wb["testcases"].values))
    assert headers[-4:] == ["amount", "expected_amount_0", "expected_amount_1", "expected_amount_2"]
    assert wb["testcases"].max_row == 1 and "settings" not in wb.sheetnames
    assert wb["steps"].cell(2, 12).value == "items"
    assert header_result_blocks([s.model_dump() for s in generated.steps], headers) == 3
    # User's existing data and settings must survive preparation with extra empty columns.
    settings = wb.create_sheet("settings")
    for row in [("key", "value"), ("url", "https://example.test"), ("screen_flow", "entry,result"), ("result_screen", "result")]:
        settings.append(row)
    wb["testcases"].append(["USER_CASE", "synthetic", "N", "RM", "X", "10;20;30", "10", "20", "30"])
    source = io.BytesIO()
    wb.save(source)
    wb.close()
    prepared = load_workbook(io.BytesIO(merge(source.getvalue(), content)))
    assert list(prepared["testcases"].values)[1][-4:] == ("10;20;30", "10", "20", "30")
    prepared.close()
    draft_path, output = tmp_path / "draft.xlsx", tmp_path / "composed.xlsx"
    draft_path.write_bytes(content)
    compose([draft_path], output)
    wb = load_workbook(output)
    assert list(next(wb["testcases"].values)) == headers
    assert wb["testcases"].max_row == 1
    wb.close()


@pytest.mark.parametrize("change", [
    {"action": "wait", "value_source": "empty", "wait_selector": ":not(*)"},
    {"action": "force_fill"}, {"value_source": "account", "step": "username"},
    {"group": "invalid group"},
])
def test_repeat_plan_rejects_semantics_not_supported_by_repeat_executor(change):
    data = plan().model_dump()
    data["steps"][0].update(change)
    with pytest.raises(ValueError):
        WorkbookPlan.model_validate(data)


def test_repeat_group_requires_values_single_screen_and_contiguous_steps():
    original = plan().model_dump()
    variants = [
        {"steps": [original["steps"][1]]},
        {"steps": [original["steps"][0], {**original["steps"][1], "screen": "other"}]},
        {"steps": [original["steps"][0], original["steps"][2], original["steps"][1]]},
    ]
    for data in variants:
        with pytest.raises(ValueError):
            WorkbookPlan.model_validate(data)


def test_header_count_limits_collisions_and_partial_header_rejected():
    for count in (0, 101, True):
        with pytest.raises(ValueError):
            planned_workbook(plan(), result_blocks=count)
    data = plan().model_dump()
    data["steps"][0]["step"] = "expected_amount_2"
    with pytest.raises(ValueError, match="collides"):
        planned_workbook(WorkbookPlan.model_validate(data), result_blocks=3)
    with pytest.raises(ValueError):
        header_result_blocks([s.model_dump() for s in plan().steps], ["expected_amount_2"])


def test_actual_repeat_executor_resolves_each_reference_without_replaying_empty_values(monkeypatch):
    path = Path(__file__).resolve().parents[2] / "docs" / "runner.py"
    spec = importlib.util.spec_from_file_location("repeat_runner_test", path)
    runner = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(runner)
    import pandas as pd
    steps = pd.DataFrame([{"step": "amount", "action": "fill", "locator_type": "css",
        "locator": "input-{i}", "value_source": "testcase", "wait_selector": ""}])
    calls = []
    monkeypatch.setattr(runner.os, "getenv", lambda name: "FAKE" if name == "SYNTHETIC_AMOUNT" else None)
    monkeypatch.setattr(runner, "do_fill", lambda page, kind, selector, value: calls.append((selector, value)))
    runner.run_repeat_group(object(), "items", steps, {"amount": "${SYNTHETIC_AMOUNT};;literal"}, {})
    assert calls == [("input-0", "FAKE"), ("input-2", "literal")]
    calls.clear()
    with pytest.raises(ValueError):
        runner.run_repeat_group(object(), "items", steps, {"amount": "${MISSING_SYNTHETIC}"}, {})
    assert calls == []


def test_preflight_blocks_repeat_actions_the_executor_would_skip_or_misinterpret():
    from runner_agent.preflight import check
    from src.runner.preflight_contract import metadata
    wb = load_workbook(io.BytesIO(planned_workbook(plan())))
    wb["steps"].cell(2, 3, "force_fill")
    wb["steps"].cell(2, 7, "Y")
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    report = metadata(check(output.getvalue()))
    assert {"sheet": "steps", "row": 2, "code": "INVALID_REPEAT_GROUP"} in report["issues"]
