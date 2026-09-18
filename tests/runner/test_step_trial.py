import io
import json
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

from runner_agent.authoring import Recording, draft_workbook
from try_step_runner import try_step, select_step


def source(tmp_path, action="click"):
    recording = Recording()
    recording.accept({"action": "click", "locator": "button:nth-of-type(1)"})
    wb = load_workbook(io.BytesIO(draft_workbook(recording)))
    wb["steps"].cell(2, 3, action)
    if action == "wait":
        wb["steps"].cell(2, 8, "button:nth-of-type(1)")
    if action in {"fill", "select"}:
        wb["steps"].cell(2, 6, "testcase")
    path = tmp_path / "source.xlsx"
    wb.save(path)
    wb.close()
    return path


def browser_mock(monkeypatch, report, fail=False):
    import playwright.sync_api
    calls, closed = [], []
    state = {"visible": True, "identity": True, "count": 1}
    class Element:
        def is_visible(self): return state["visible"]
        def act(self, name, *args, **kwargs):
            intent = json.loads(report.read_text())
            assert intent["status"] == "EXECUTION_STARTED" and intent["attempts"] == 1
            calls.append((name, args, kwargs))
            if fail:
                raise RuntimeError("synthetic-private-value")
        def click(self, **kwargs): self.act("click", **kwargs)
        def check(self, **kwargs): self.act("check", **kwargs)
        def fill(self, value, **kwargs): self.act("fill", value, **kwargs)
        def select_option(self, **kwargs): self.act("select", **kwargs)
        def wait_for_element_state(self, value, **kwargs): self.act("wait", value, **kwargs)
    element = Element()
    locator = SimpleNamespace(count=lambda: state["count"], is_visible=element.is_visible,
        highlight=lambda: None, element_handle=lambda: element,
        evaluate=lambda script, handle: state["identity"] and handle is element)
    page = object()
    context = SimpleNamespace(new_page=lambda: page, pages=[page])
    browser = SimpleNamespace(new_context=lambda: context, is_connected=lambda: True, close=lambda: closed.append(True))
    class Playwright:
        def __enter__(self): return SimpleNamespace(chromium=SimpleNamespace(launch=lambda **kw: browser))
        def __exit__(self, *args): pass
    monkeypatch.setattr(playwright.sync_api, "sync_playwright", Playwright)
    return locator, state, calls, closed


@pytest.mark.parametrize("action", ["fill", "click", "select", "check", "wait"])
def test_try_exactly_one_reviewed_element_and_persist_intent_before_action(tmp_path, monkeypatch, action):
    config, report = source(tmp_path, action), tmp_path / "trial.json"
    original = config.read_bytes()
    locator, _, calls, closed = browser_mock(monkeypatch, report)
    answers = iter(["1", "EXECUTE 2"])
    result = try_step(config, 2, report, prompt=lambda _: next(answers),
        value_prompt=lambda _: "synthetic-private-value", builder=lambda *a: locator)
    assert result["status"] == "ACTION_COMPLETED" and result["attempts"] == 1
    assert len(calls) == 1 and calls[0][0] == action and closed == [True]
    assert config.read_bytes() == original
    assert "synthetic-private" not in report.read_text()
    assert "selector" not in report.read_text() and "locator" not in report.read_text()
    with pytest.raises(ValueError):
        try_step(config, 2, report)
    assert len(calls) == 1


@pytest.mark.parametrize("case,expected", [("cancel", "CANCELLED"), ("changed", "WORKBOOK_CHANGED"),
    ("identity", "TARGET_CHANGED"), ("ambiguous", "TARGET_UNAVAILABLE"), ("hidden", "TARGET_UNAVAILABLE")])
def test_trial_does_not_act_without_confirmed_unchanged_source_and_target(tmp_path, monkeypatch, case, expected):
    config, report = source(tmp_path), tmp_path / "trial.json"
    locator, state, calls, _ = browser_mock(monkeypatch, report)
    if case == "ambiguous": state["count"] = 2
    if case == "hidden": state["visible"] = False
    answers = iter(["1", "cancel" if case == "cancel" else "EXECUTE 2"])
    def prompt(_):
        answer = next(answers)
        if answer.startswith("EXECUTE"):
            if case == "changed": config.write_bytes(config.read_bytes() + b"changed")
            if case == "identity": state["identity"] = False
        return answer
    result = try_step(config, 2, report, prompt=prompt, builder=lambda *a: locator)
    assert result["status"] == expected and calls == []
    assert json.loads(report.read_text())["attempts"] == 0


def test_failure_is_unknown_not_failed_testcase_and_never_retried(tmp_path, monkeypatch, capsys):
    config, report = source(tmp_path), tmp_path / "trial.json"
    locator, _, calls, closed = browser_mock(monkeypatch, report, fail=True)
    answers = iter(["1", "EXECUTE 2"])
    result = try_step(config, 2, report, prompt=lambda _: next(answers), builder=lambda *a: locator)
    assert result["status"] == "OUTCOME_UNKNOWN" and len(calls) == 1 and closed == [True]
    assert "synthetic-private" not in capsys.readouterr().out + report.read_text()


def test_cannot_act_when_persisting_intent_fails(tmp_path, monkeypatch):
    import try_step_runner
    config, report = source(tmp_path), tmp_path / "trial.json"
    locator, _, calls, _ = browser_mock(monkeypatch, report)
    real_write = try_step_runner.write_report
    def write(path, data, initial=False):
        if data["status"] == "EXECUTION_STARTED":
            raise OSError("synthetic storage failure")
        return real_write(path, data, initial)
    monkeypatch.setattr(try_step_runner, "write_report", write)
    answers = iter(["1", "EXECUTE 2"])
    result = try_step(config, 2, report, prompt=lambda _: next(answers), builder=lambda *a: locator)
    assert result["status"] == "OUTCOME_UNKNOWN" and calls == []


@pytest.mark.parametrize("changes", [{"group": "items"}, {"action": "force_fill"}, {"action": "read_result_single"},
    {"wait_selector": "timeout_reload:1"}, {"prefill_check": "Y"}, {"locator": ":not(*)"},
    {"locator": "input-{i}"}, {"locator_type": "form_item"}])
def test_complex_and_unresolved_steps_need_full_runner(changes):
    step = {"row": 2, "action": "click", "locator": "button", "locator_type": "css", **changes}
    with pytest.raises(ValueError):
        select_step({"steps": [step]}, 2)
