import importlib.util
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("uat_runner", ROOT / "docs" / "runner.py")
runner = importlib.util.module_from_spec(spec)
spec.loader.exec_module(runner)


@pytest.mark.parametrize("threshold,expected", [("exact", False), ("±0.1", True),
                                                ("5%", True), ("±1e-3", False)])
def test_numeric_comparison(threshold, expected):
    assert runner._compare_values("2", "2.05", threshold)[0] is expected


def test_threshold_cannot_execute_python(tmp_path):
    marker = tmp_path / "executed"
    with pytest.raises(ValueError):
        runner._compare_values("1", "1", f"±__import__('pathlib').Path({str(marker)!r}).touch()")
    assert not marker.exists()


def test_input_values_never_printed(capsys, monkeypatch):
    monkeypatch.setattr(runner, "do_fill", lambda *a, **k: None)
    step = pd.Series({"action": "fill", "locator_type": "label", "locator": "User",
                      "wait_selector": "", "step": "username"})
    runner.run_step(None, step, "synthetic-user-value")
    assert "synthetic-user-value" not in capsys.readouterr().out


def test_validator_rejects_unimplemented_expect():
    steps = pd.DataFrame([{"screen": "login", "step": "check", "action": "expect",
                           "locator_type": "css", "locator": "h1", "value_source": "",
                           "active": "Y", "wait_selector": ""}])
    assert not runner.validate_config(steps, [], {"url": "https://example.test", "screen_flow": "login"})


class Page:
    def goto(self, *a, **k): pass
    def wait_for_load_state(self, *a, **k): pass


@pytest.mark.parametrize("raises,expected", [(False, "UNVERIFIED"), (True, "ERROR")])
def test_case_without_assertions_is_never_passed(tmp_path, monkeypatch, raises, expected):
    monkeypatch.setattr(runner, "get_account", lambda *a: {"role_code": "RM", "specialized_bank": "X"})
    def screen(*a, **k):
        if raises:
            raise RuntimeError("synthetic private input")
    monkeypatch.setattr(runner, "run_screen", screen)
    runner._case_outcomes = []
    runner.run_testcase(Page(), {"tc_id": "T1", "mo_ta": "case", "role_code": "RM", "specialized_bank": "X"},
                        pd.DataFrame(columns=["screen", "active"]),
                        {"url": "https://example.test", "screen_flow": "login", "screenshot_on_error": "N"},
                        {"prefix": "run", "errors": tmp_path}, tmp_path / "log.txt", is_first=True)
    assert runner._case_outcomes[0]["status"] == expected
    assert "synthetic private input" not in (tmp_path / "log.txt").read_text()


def test_export_redacts_registered_values(tmp_path):
    from openpyxl import load_workbook
    runner._redactor.register("synthetic-sensitive")
    target = tmp_path / "results.xlsx"
    runner.export_results_simple([{"tc_id": "T1", "overall": "PASS", "real_name": "synthetic-sensitive"}], str(target))
    wb = load_workbook(target)
    assert "synthetic-sensitive" not in str(list(wb.active.values))
    wb.close()


def test_export_excel_single_rows_do_not_create_extra_blocks(tmp_path):
    """Row read_result_single không có key "block": không được thành b1/b2 (block chấm điểm thừa)."""
    from openpyxl import load_workbook
    group = lambda i, ok: {"tc_id": "TC1", "mo_ta": "m", "id_ho_so": "H", "overall": "PASS" if ok else "FAIL",
                           "block": i, "exp_gia": "1", "real_gia": "1", "pass_gia": "PASS" if ok else "FAIL"}
    single = {"tc_id": "TC1", "mo_ta": "m", "id_ho_so": "H", "overall": "PASS",
              "exp_tong": "9", "real_tong": "9", "pass_tong": "PASS"}
    out = tmp_path / "r.xlsx"
    runner.export_results_excel([group(0, True), group(1, True), single], str(out))
    headers = [c.value for c in next(load_workbook(out)["results"].iter_rows(min_row=1, max_row=1))]
    assert not any(h.startswith("b2_") for h in headers)
    assert any(h.startswith("b1_") for h in headers) and not any("tong" in h for h in headers)
    # overall vẫn tính cả row single
    failed = [group(0, True), group(1, True), {**single, "overall": "FAIL"}]
    runner.export_results_excel(failed, str(out))
    ws = load_workbook(out)["results"]
    assert [c.value for c in ws[2]][3] == "FAIL"
