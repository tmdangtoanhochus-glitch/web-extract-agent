import re
import os
import sys
import argparse
import threading
from pathlib import Path
from datetime import datetime
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
import pandas as pd
from dotenv import load_dotenv
from playwright.sync_api import sync_playwright, Page
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from runner_agent.sanitize import Redactor
from src.runner.scoped_locator import PREFIX as SCOPE_PREFIX, resolve as resolve_scope

_redactor = Redactor()
_case_outcomes = []

# ── Base directory ───────────────────────────────────────

def get_base_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


# ── Constants ────────────────────────────────────────────
SPINNER_TIMEOUT       = 120000
POPUP_TIMEOUT         = 5000
SCREENSHOT_DIR        = "errors"
POPUP_BUTTON_LOCATOR  = ".ant-modal-footer button.ant-btn-primary"
DEBUG                 = False
BLOCK_TRIGGER_SEL     = r"^Kết quả chấm điểm"
RESULT_WAIT_SELECTOR  = "[id^='forms_SCORING_RESULT_elements_']"
ID_FIELD_LABEL        = "ID hồ sơ"
ID_FIELD_VALUE_SEL    = "span.text-textPrimary"
RESULT_SCREEN         = "scoring_result"
SPINNER_STROKE_COLOR  = "#F4600C"
STEP_WRAPPER_SELECTOR = ".z-20"
SPINNER_WRAPPER_CLASS = "bg-warning50"
TICK_CLASS            = "text-primaryColor"
STEP_LINK_XPATH       = "following-sibling::div[contains(@class,'flex-col')][1]//a"
HOME_BUTTON_SELECTOR  = "img.ant-image-img.cursor-pointer"
LOGIN_SCREEN          = "login"
REUSE_SESSION         = True
_print_lock           = threading.Lock()
_log_lock             = threading.Lock()


# ── Thread-safe helpers ───────────────────────────────────

def tprint(*args, tc_id: str = "", **kwargs):
    """Print thread-safe, them [tc_id] prefix khi chay song song."""
    with _print_lock:
        if tc_id:
            msg = _redactor.text(" ".join(str(a) for a in args))
            print(f"[{_redactor.text(tc_id)}] {msg}", **kwargs)
        else:
            print(*[_redactor.text(a) for a in args], **kwargs)


# ── CLI & Config path ────────────────────────────────────

def parse_args():
    parser = argparse.ArgumentParser(description="UAT Automation Runner")
    parser.add_argument("--config", default="", help="Duong dan file config xlsx")
    return parser.parse_args()


def pick_config() -> Path:
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk(); root.withdraw()
        path = filedialog.askopenfilename(
            title="Chon file config UAT",
            initialdir=str(get_base_dir() / "config"),
            filetypes=[("Excel files", "*.xlsx")]
        )
        root.destroy()
        return Path(path) if path else None
    except Exception:
        p = input("Nhap duong dan file config: ").strip()
        return Path(p) if p else None


def get_output_dirs(config_path: Path) -> dict:
    folder = config_path.stem
    base = get_base_dir() / "output" / folder
    dirs = {
        "base":    base,
        "logs":    base / "run_log",
        "results": base / "results",
        "errors":  base / "errors",
        "prefix":  folder,
    }
    for d in ["logs", "results", "errors"]:
        dirs[d].mkdir(parents=True, exist_ok=True)
    return dirs


# ── Load config ──────────────────────────────────────────

def load_settings(path: Path) -> dict:
    df = pd.read_excel(path, sheet_name="settings", dtype=str).fillna("")
    return dict(zip(df["key"], df["value"]))


def load_steps(path: Path, sheet_name: str = "steps") -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet_name, dtype=str).fillna("")


def load_testcases(path: Path) -> list[dict]:
    from openpyxl import load_workbook
    wb   = load_workbook(path, data_only=True)
    ws   = wb["testcases"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h) if h is not None else "" for h in rows[0]]
    result  = []
    for row in rows[1:]:
        tc = {}
        for h, v in zip(headers, row):
            if v is None:
                tc[h] = ""
            elif hasattr(v, "strftime"):
                tc[h] = v.strftime("%d/%m/%Y")
            else:
                tc[h] = str(v).strip()
        if tc.get("active", "").upper() == "Y":
            result.append(tc)
    return result


def filter_testcases(testcases: list[dict], settings: dict) -> list[dict]:
    run_filter = settings.get("run_filter", "").strip()
    if not run_filter:
        return testcases
    ids = [tc["tc_id"] for tc in testcases]
    if "," in run_filter:
        only_ids = {x.strip() for x in run_filter.split(",")}
        result = [tc for tc in testcases if tc["tc_id"] in only_ids]
        tprint(f"  [Filter] run_only: {only_ids}")
        return result
    if ":" in run_filter:
        parts = run_filter.split(":")
        if len(parts) == 2:
            start, end = parts[0].strip(), parts[1].strip()
            if start in ids and end in ids:
                s_idx = ids.index(start)
                e_idx = ids.index(end)
                result = testcases[s_idx:e_idx+1]
                tprint(f"  [Filter] run_range: {start} → {end} ({len(result)} TC)")
                return result
    if run_filter in ids:
        result = testcases[ids.index(run_filter):]
        tprint(f"  [Filter] start_from: {run_filter} ({len(result)} TC)")
        return result
    tprint(f"  [WARN] run_filter '{run_filter}' khong hop le, chay tat ca")
    return testcases


def get_account(role_code: str, specialized_bank: str) -> dict:
    if role_code and specialized_bank:
        prefix = f"{role_code.upper()}_{specialized_bank.upper()}"
    elif role_code:
        prefix = role_code.upper()
    else:
        prefix = "DEFAULT"
    username = os.getenv(f"{prefix}_USERNAME")
    password = os.getenv(f"{prefix}_PASSWORD")
    if not username or not password:
        raise ValueError(
            f"Khong tim thay account trong .env: "
            f"{prefix}_USERNAME / {prefix}_PASSWORD"
        )
    _redactor.register(username, password)
    return {
        "username":         username,
        "password":         password,
        "role_code":        role_code,
        "specialized_bank": specialized_bank,
    }


# ── Ghi log ──────────────────────────────────────────────

def write_log(log_path: Path, tc_id: str, mo_ta: str, status: str):
    if not log_path:
        return
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with _log_lock:
        with open(log_path, "a", encoding="utf-8") as f:
            f.write(_redactor.text(f"{ts} | {status:<6} | {tc_id} | {mo_ta}\n"))


# ── Locator builder ──────────────────────────────────────

def build_locator(page: Page, locator_type: str, locator: str):
    if locator_type.lower() in {"css", ""} and locator.startswith(SCOPE_PREFIX):
        return resolve_scope(page, locator)
    match locator_type.lower():
        case "css":
            return page.locator(locator)
        case "xpath":
            return page.locator(f"xpath={locator}")
        case "role":
            role, name = locator.split("|", 1)
            return page.get_by_role(role, name=name)
        case "text":
            return page.get_by_text(locator, exact=True)
        case "label":
            return page.get_by_label(locator)
        case "filter":
            parts = locator.split("|")
            text  = parts[1]
            has_text_filter = text[1:] if text.startswith("~") else re.compile(rf"^{re.escape(text)}$")
            result = (
                page.locator(f"[class='{parts[0]}']")
                .filter(has_text=has_text_filter)
            )
            if len(parts) > 2:
                result = result.locator(parts[2])
            return result
        case "form_item":
            return (
                page.locator(".ant-form-item")
                .filter(has_text=locator)
                .locator("input.ant-input")
            )
        case "nth":
            parts = locator.split("|")
            return page.locator(parts[0]).nth(int(parts[1]))
        case "role_nth":
            parts = locator.split("|")
            return page.get_by_role(parts[0], name=parts[1]).nth(int(parts[2]))
        case _:
            return page.locator(locator)


# ── _reload_and_navigate ─────────────────────────────────

def _reload_and_navigate(page: Page, tc_id: str = ""):
    page.reload(wait_until="domcontentloaded", timeout=60000)
    tprint("      [reload] Trang da reload", tc_id=tc_id)

    try:
        page.wait_for_selector(f"{STEP_WRAPPER_SELECTOR} > svg", timeout=15000)
        page.wait_for_timeout(1000)
    except Exception as e:
        tprint(f"      [reload] Khong tim thay danh sach buoc: {e}", tc_id=tc_id)
        return False

    try:
        all_steps = page.locator(STEP_WRAPPER_SELECTOR)
        total     = all_steps.count()
        if DEBUG:
            tprint(f"      [DEBUG] Tong so buoc: {total}", tc_id=tc_id)

        spinner_targets = []
        tick_targets    = []

        for idx in range(total):
            div = all_steps.nth(idx)
            try:
                svg = div.locator(":scope > svg").first
                if svg.count() == 0:
                    continue

                svg_cls      = svg.get_attribute("class") or ""
                orange_paths = svg.locator(f"path[stroke='{SPINNER_STROKE_COLOR}']").count()
                clip_groups  = svg.locator("g[clip-path]").count()

                is_tick    = TICK_CLASS in svg_cls
                is_spinner = not is_tick and (orange_paths >= 4 or clip_groups > 0)
                icon_type  = "TICK" if is_tick else "SPINNER" if is_spinner else "OTHER"

                link = div.locator(f"xpath={STEP_LINK_XPATH}").first
                text = (
                    re.sub(r"\s+", " ", link.inner_text()).strip()
                    if link.count() > 0 else "NO_TEXT"
                )
                if DEBUG:
                    tprint(f"      [DEBUG] [{idx}] {icon_type} | {text[:50]}", tc_id=tc_id)

                if link.count() == 0:
                    continue

                item = {"link": link, "text": text}
                if is_spinner:
                    spinner_targets.append(item)
                elif is_tick:
                    tick_targets.append(item)

            except Exception as ex:
                if DEBUG:
                    tprint(f"      [DEBUG] [{idx}] loi: {ex}", tc_id=tc_id)

        if DEBUG:
            tprint(f"      [DEBUG] spinner={len(spinner_targets)} tick={len(tick_targets)}", tc_id=tc_id)

        if spinner_targets:
            target = spinner_targets[-1]
            if DEBUG:
                tprint("      [DEBUG] Chon spinner", tc_id=tc_id)
        elif tick_targets:
            target = tick_targets[-1]
            if DEBUG:
                tprint("      [DEBUG] Fallback tick cuoi", tc_id=tc_id)
        else:
            tprint("      [reload] Khong tim thay buoc nao", tc_id=tc_id)
            return False

        link = target["link"]
        text = target["text"]
        if DEBUG:
            tprint(f"      [DEBUG] Navigate den: {text[:50]}", tc_id=tc_id)
        link.wait_for(state="visible", timeout=10000)
        link.scroll_into_view_if_needed()
        link.click()
        page.wait_for_timeout(1500)
        if DEBUG:
            tprint(f"      [DEBUG] Da click, URL: {page.url}", tc_id=tc_id)
        return True

    except Exception as e:
        tprint(f"      [reload] Loi khi navigate: {e}", tc_id=tc_id)
        return False


# ── Action handlers ──────────────────────────────────────

def do_fill(page: Page, lt: str, loc: str, value: str, force: bool = False):
    el = build_locator(page, lt, loc)
    el.fill(value, force=True) if force else el.fill(value)


def do_fill_enter(page: Page, lt: str, loc: str, value: str):
    el = build_locator(page, lt, loc)
    el.click(); el.fill(""); el.type(value); el.press("Enter")


def do_click(page: Page, lt: str, loc: str, wait_sel: str, tc_id: str = ""):
    if loc.startswith(SCOPE_PREFIX):
        build_locator(page, lt, loc).click()
        if wait_sel:
            do_wait(page, wait_sel, tc_id=tc_id)
        return
    build_locator(page, lt, loc).click()
    if not wait_sel:
        return

    if wait_sel.startswith("hidden_or_popup:"):
        import time
        sel   = wait_sel[len("hidden_or_popup:"):]
        start = time.monotonic()
        iteration = 0
        spinner_appeared = False
        try:
            page.wait_for_selector(sel, state="visible", timeout=3000)
            spinner_appeared = True
            tprint("      [do_click] Spinner xuat hien, bat dau poll...", tc_id=tc_id)
        except Exception:
            tprint("      [do_click] Spinner khong xuat hien → tiep tuc", tc_id=tc_id)

        if not spinner_appeared:
            return

        while (time.monotonic() - start) * 1000 < SPINNER_TIMEOUT:
            iteration += 1
            spinner_visible = False
            try:
                spinner_visible = page.locator(sel).first.is_visible()
            except Exception:
                pass
            if not spinner_visible:
                tprint(f"      [do_click] Spinner tat sau {iteration} lan check", tc_id=tc_id)
                break
            try:
                tiep_tuc    = page.locator(POPUP_BUTTON_LOCATOR)
                popup_visible = tiep_tuc.first.is_visible()
                elapsed_ms  = (time.monotonic() - start) * 1000

                if DEBUG and (iteration <= 3 or popup_visible or iteration % 10 == 0):
                    tprint(
                        f"      [DEBUG] "
                        f"iter={iteration} | "
                        f"elapsed={elapsed_ms/1000:.1f}s | "
                        f"limit={SPINNER_TIMEOUT/1000:.1f}s | "
                        f"spinner={spinner_visible} | "
                        f"popup={popup_visible} | "
                        f"selector='{POPUP_BUTTON_LOCATOR}'",
                        tc_id=tc_id
                    )
                if popup_visible:
                    try:
                        tiep_tuc.first.click(force=True, timeout=3000)
                        tprint("      [do_click] Da click popup", tc_id=tc_id)
                    except Exception:
                        try:
                            tiep_tuc.first.dispatch_event("click")
                        except Exception:
                            pass
                    tprint("      [do_click] Cho spinner tat sau khi click popup...", tc_id=tc_id)
                    try:
                        page.wait_for_selector(sel, state="hidden", timeout=POPUP_TIMEOUT)
                        tprint("      [do_click] Spinner da tat, tiep tuc", tc_id=tc_id)
                    except Exception:
                        tprint("      [do_click] Spinner van con → reload", tc_id=tc_id)
                        _reload_and_navigate(page, tc_id=tc_id)
                    break
            except Exception as e:
                if DEBUG:
                    tprint(f"      [DEBUG] iter={iteration} loi: {e}", tc_id=tc_id)
            page.wait_for_timeout(500)
        else:
            tprint("      [do_click] Spinner timeout → reload", tc_id=tc_id)
            _reload_and_navigate(page, tc_id=tc_id)

    elif wait_sel.startswith("hidden:"):
        sel = wait_sel[len("hidden:"):]
        try:
            page.locator(sel).first.wait_for(state="hidden", timeout=SPINNER_TIMEOUT)
        except Exception:
            tprint("      [do_click] Spinner timeout → reload", tc_id=tc_id)
            _reload_and_navigate(page, tc_id=tc_id)

    elif wait_sel.startswith("wait_or_reload:"):
        sel = wait_sel[len("wait_or_reload:"):]
        try:
            page.wait_for_selector(sel, timeout=POPUP_TIMEOUT)
        except Exception:
            tprint("      [do_click] Element chua xuat hien → reload", tc_id=tc_id)
            _reload_and_navigate(page, tc_id=tc_id)

    elif wait_sel.startswith("timeout_reload:"):
        ms = int(wait_sel[len("timeout_reload:"):])
        page.wait_for_timeout(ms)
        _reload_and_navigate(page, tc_id=tc_id)

    elif wait_sel.startswith("timeout:"):
        page.wait_for_timeout(int(wait_sel[len("timeout:"):]))

    elif wait_sel == "spin":
        page.locator(".ant-spin-spinning").first.wait_for(
            state="hidden", timeout=SPINNER_TIMEOUT)

    else:
        page.wait_for_selector(wait_sel, state="visible")


def do_click_if_exists(page: Page, lt: str, loc: str, wait_sel: str,
                        tc_id: str = ""):
    if not wait_sel or not wait_sel.startswith("hidden:"):
        if wait_sel and wait_sel.startswith("timeout:"):
            page.wait_for_timeout(int(wait_sel[len("timeout:"):]))
        elif wait_sel:
            page.wait_for_selector(wait_sel, state="visible")
        return

    sel = wait_sel[len("hidden:"):]

    try:
        page.wait_for_selector(sel, state="visible", timeout=3000)
        tprint("      [click_if_exists] Spinner xuat hien, dang cho tat...", tc_id=tc_id)
    except Exception:
        tprint("      [click_if_exists] Spinner chua xuat hien → tiep tuc", tc_id=tc_id)
        return

    try:
        page.wait_for_selector(sel, state="hidden", timeout=POPUP_TIMEOUT)
        tprint("      [click_if_exists] Spinner da tat, tiep tuc", tc_id=tc_id)
        return
    except Exception:
        pass

    tprint(f"      [click_if_exists] Spinner van con sau {POPUP_TIMEOUT}ms → reload", tc_id=tc_id)
    if DEBUG:
        tprint(f"      [DEBUG] URL: {page.url}", tc_id=tc_id)
        try:
            btns = [b for b in page.get_by_role("button").all() if b.is_visible()]
            tprint(f"      [DEBUG] Buttons: {[b.inner_text().strip()[:30] for b in btns]}", tc_id=tc_id)
            popup_visible = page.locator(POPUP_BUTTON_LOCATOR).count() > 0
            tprint(f"      [DEBUG] Popup visible: {popup_visible}", tc_id=tc_id)
        except Exception as de:
            tprint(f"      [DEBUG] Loi: {de}", tc_id=tc_id)

    try:
        tiep_tuc = build_locator(page, lt, loc) if lt and loc else page.locator(POPUP_BUTTON_LOCATOR)
        tiep_tuc.first.wait_for(state="visible", timeout=2000)
        tiep_tuc.first.click()
        tprint("      [click_if_exists] Da click popup", tc_id=tc_id)
    except Exception:
        tprint("      [click_if_exists] Khong co popup", tc_id=tc_id)

    _reload_and_navigate(page, tc_id=tc_id)


def do_check(page: Page, lt: str, loc: str):
    build_locator(page, lt, loc).check()


def do_radio(page: Page, lt: str, loc: str, value: str):
    page.locator(loc).get_by_text(value, exact=True).click()


def _get_dropdown_selector(dropdown_sel: str) -> str:
    if dropdown_sel == "not_hidden":
        return ".ant-select-dropdown:not(.ant-select-dropdown-hidden)"
    return ".ant-select-dropdown:visible"


def _get_option_filter(value: str, match_type: str):
    if match_type == "exact":
        return re.compile(rf"^{re.escape(value)}$")
    return value


def _wait_dropdown_close(page: Page, dropdown_selector: str):
    try:
        page.wait_for_selector(dropdown_selector, state="hidden", timeout=3000)
    except Exception:
        pass


def do_select_antd(page: Page, lt: str, loc: str, value: str,
                   dropdown_sel: str = "visible", match_type: str = "contains"):
    input_el   = build_locator(page, lt, loc)
    dd_sel     = _get_dropdown_selector(dropdown_sel)
    opt_filter = _get_option_filter(value, match_type)
    input_el.click(); input_el.fill(value)
    option = (
        page.locator(dd_sel).locator(".ant-select-item-option")
        .filter(has_text=opt_filter).first
    )
    option.wait_for(state="visible"); option.click()
    _wait_dropdown_close(page, dd_sel)


def do_force_select_antd(page: Page, lt: str, loc: str, value: str,
                          dropdown_sel: str = "not_hidden",
                          match_type: str = "exact"):
    input_el   = build_locator(page, lt, loc)
    dd_sel     = _get_dropdown_selector(dropdown_sel)
    opt_filter = _get_option_filter(value, match_type)
    input_el.fill(value, force=True)
    option = (
        page.locator(dd_sel).locator(".ant-select-item-option")
        .filter(has_text=opt_filter).first
    )
    option.wait_for(state="visible"); option.click()
    _wait_dropdown_close(page, dd_sel)


def do_wait(page: Page, wait_sel: str, tc_id: str = ""):
    if wait_sel.startswith(SCOPE_PREFIX):
        resolve_scope(page, wait_sel).wait_for(state="visible")
        return
    if wait_sel == "spin":
        page.locator(".ant-spin-spinning").first.wait_for(
            state="hidden", timeout=SPINNER_TIMEOUT)
    elif wait_sel.startswith("hidden:"):
        sel = wait_sel[len("hidden:"):]
        page.locator(sel).first.wait_for(state="hidden", timeout=SPINNER_TIMEOUT)
    elif wait_sel.startswith("wait_or_reload:"):
        sel = wait_sel[len("wait_or_reload:"):]
        try:
            page.wait_for_selector(sel, timeout=POPUP_TIMEOUT)
        except Exception:
            tprint("    [wait] Element chua xuat hien → reload", tc_id=tc_id)
            _reload_and_navigate(page, tc_id=tc_id)
    elif wait_sel.startswith("timeout_reload:"):
        ms = int(wait_sel[len("timeout_reload:"):])
        tprint(f"    [wait] Cho {ms}ms roi reload", tc_id=tc_id)
        page.wait_for_timeout(ms)
        _reload_and_navigate(page, tc_id=tc_id)
    elif wait_sel.startswith("timeout:"):
        page.wait_for_timeout(int(wait_sel[len("timeout:"):]))
    else:
        page.wait_for_selector(wait_sel, state="visible")


def do_select(page: Page, lt: str, loc: str, value: str):
    build_locator(page, lt, loc).select_option(label=value)


def do_fill_sequence(page: Page, locators_str: str, values_str: str):
    labels = [l.strip() for l in locators_str.split("|")]
    values = [v.strip() for v in values_str.split(";")]
    for i, label in enumerate(labels):
        value = values[i] if i < len(values) else ""
        if not value:
            continue
        (
            page.locator(".ant-form-item")
            .filter(has_text=label)
            .locator("input.ant-input")
            .fill(value)
        )


# ── Read result ───────────────────────────────────────────

def read_result_field(page: Page, step: pd.Series, block_idx: int,
                       tc_id: str = "") -> str:
    method  = step.get("read_method", "").strip().lower()
    locator = step["locator"].replace("{i}", str(block_idx))
    try:
        match method:
            case "label_input":
                return (
                    page.locator(f"label[for='{locator}']")
                    .locator("xpath=ancestor::div[contains(@class,'ant-form-item')]")
                    .locator("input").first.input_value() or ""
                )
            case "css_input":
                return build_locator(page, "css", locator).input_value() or ""
            case "label_span_title":
                return (
                    page.locator(f"label[for='{locator}']")
                    .locator("xpath=ancestor::div[contains(@class,'ant-form-item')]")
                    .locator("span.ant-select-selection-item")
                    .get_attribute("title") or ""
                )
            case "button_regex":
                btn  = page.get_by_role("button", name=re.compile(BLOCK_TRIGGER_SEL)).nth(block_idx)
                text = btn.inner_text()
                m    = re.search(locator, text, re.S)
                raw  = m.group(1).strip() if m else ""
                return re.sub(r'[•\s]+$', '', raw).strip()
            case "sibling_span":
                return (
                    page.locator("span", has_text=locator)
                    .locator("xpath=..")
                    .locator("span.text-textPrimary")
                    .inner_text().strip()
                )
            case "css_disabled":
                return build_locator(page, "css", locator).evaluate("el => el.value") or ""
            case _:
                tprint(f"    [WARN] read_method '{method}' chua duoc ho tro", tc_id=tc_id)
                return ""
    except Exception as e:
        tprint(
            f"    [WARN] Khong doc duoc field method='{method}' "
            f"locator='{locator}' block={block_idx}: {e}",
            tc_id=tc_id
        )
        return ""


# ── Compare values ────────────────────────────────────────

def _is_numeric(value: str) -> bool:
    try:
        float(value.replace(",","").replace("%","").replace(" ",""))
        return True
    except ValueError:
        return False


def _compare_values(expected: str, actual: str, threshold: str) -> tuple[bool, str]:
    from difflib import SequenceMatcher
    exp = expected.strip()
    act = actual.strip()
    if threshold == "exact":
        return exp == act, "exact"
    if threshold == "contains":
        return exp.lower() in act.lower(), "contains"
    if _is_numeric(exp) and _is_numeric(act):
        exp_f = float(exp.replace(",","").replace("%",""))
        act_f = float(act.replace(",","").replace("%",""))
        if threshold.startswith("±") or threshold.startswith("+/-"):
            n_str = threshold.replace("±","").replace("+/-","").strip()
            n_str = n_str.replace("**-","e-").replace("**","e")
            n_str = re.sub(r"10e(-?\d+)", r"1e\1", n_str)
            # Config upload không được thực thi Python. Chỉ nhận số hữu hạn.
            import math
            n = float(n_str)
            if not math.isfinite(n) or n < 0:
                raise ValueError("Sai so phai la so huu han khong am")
            diff = abs(exp_f - act_f)
            return diff <= n, f"|{exp_f}-{act_f}|={diff:.2e} <= ±{n:.2e}"
        if threshold.endswith("%"):
            pct = float(threshold.rstrip("%").strip()) / 100
            if exp_f != 0:
                rel_diff = abs(exp_f - act_f) / abs(exp_f)
                return rel_diff <= pct, f"rel_diff={rel_diff:.1%} <= {threshold}"
            return act_f == 0, f"expected=0, actual={act_f}"
    try:
        rt    = float(threshold)
        ratio = SequenceMatcher(None, exp.lower(), act.lower()).ratio()
        return ratio >= rt, f"similarity={ratio:.2f} >= {rt}"
    except ValueError:
        pass
    return exp == act, "exact(fallback)"


# ── Validate config ───────────────────────────────────────

VALID_ACTIONS = {
    "fill", "force_fill", "fill_enter", "click", "click_if_exists",
    "check", "uncheck", "upload", "radio", "select_antd", "force_select_antd", "select",
    "nth", "form_item", "fill_sequence", "wait",
    "read_result", "read_result_single", "read_result_group",
}
VALID_LOCATOR_TYPES = {
    "css", "xpath", "role", "text", "label",
    "filter", "form_item", "nth", "role_nth", "",
}
VALID_VALUE_SOURCES = {"testcase", "account", "empty", "keyword", ""}
REQUIRED_STEP_COLS  = {"screen", "step", "action", "locator_type",
                        "locator", "value_source", "active", "wait_selector"}
REQUIRED_TC_COLS    = {"tc_id", "mo_ta", "active", "role_code", "specialized_bank"}
REQUIRED_SETTINGS   = {"url", "screen_flow"}


def validate_config(df_steps: pd.DataFrame, testcases: list[dict],
                    settings: dict) -> bool:
    errors   = []
    warnings = []

    for key in REQUIRED_SETTINGS:
        if not settings.get(key, "").strip():
            errors.append(f"[settings] Thieu key bat buoc: '{key}'")

    missing_cols = REQUIRED_STEP_COLS - set(df_steps.columns)
    if missing_cols:
        errors.append(f"[steps] Thieu cot: {missing_cols}")
    else:
        active_steps = df_steps[df_steps["active"].str.upper() == "Y"]
        for _, row in active_steps.iterrows():
            action = str(row.get("action", "")).strip().lower()
            if action not in VALID_ACTIONS:
                close = [a for a in VALID_ACTIONS if a.startswith(action[:4])]
                hint  = f" (y ban muon dung: {close[0]}?)" if close else ""
                errors.append(
                    f"[steps] screen='{row['screen']}' step='{row['step']}': "
                    f"action='{action}' khong hop le{hint}"
                )
            lt = str(row.get("locator_type", "")).strip().lower()
            if lt not in VALID_LOCATOR_TYPES:
                errors.append(
                    f"[steps] screen='{row['screen']}' step='{row['step']}': "
                    f"locator_type='{lt}' khong hop le"
                )
            vs = str(row.get("value_source", "")).strip().lower()
            if vs not in VALID_VALUE_SOURCES:
                warnings.append(
                    f"[steps] screen='{row['screen']}' step='{row['step']}': "
                    f"value_source='{vs}' la la"
                )

        screen_flow = settings.get("screen_flow", "")
        if screen_flow:
            flow_screens  = {s.strip() for s in screen_flow.split(",") if s.strip()}
            steps_screens = set(df_steps["screen"].unique())
            for sc in flow_screens:
                if sc not in steps_screens:
                    warnings.append(
                        f"[settings] screen_flow chua screen '{sc}' "
                        f"nhung khong co step nao trong sheet steps"
                    )

        dupes = (
            active_steps.groupby(["screen", "step"])
            .size().reset_index(name="count")
        )
        for _, d in dupes[dupes["count"] > 1].iterrows():
            warnings.append(
                f"[steps] Duplicate screen='{d['screen']}' "
                f"step='{d['step']}' ({d['count']} lan)"
            )

    if testcases:
        tc_cols = set(testcases[0].keys())
        missing_tc = REQUIRED_TC_COLS - tc_cols
        if missing_tc:
            errors.append(f"[testcases] Thieu cot: {missing_tc}")
        tc_ids = [tc.get("tc_id", "") for tc in testcases]
        if len(tc_ids) != len(set(tc_ids)):
            duped = [x for x in set(tc_ids) if tc_ids.count(x) > 1]
            errors.append(f"[testcases] Duplicate tc_id: {duped}")
        if any(not re.fullmatch(r"[\w-]{1,80}", str(t)) for t in tc_ids):
            errors.append("[testcases] tc_id chi duoc chua chu, so, _ va -")
        roles = {(t.get("role_code", ""), t.get("specialized_bank", "")) for t in testcases}
        if len(roles) > 1 and settings.get("reuse_session", "Y").upper() == "Y":
            errors.append("[settings] Nhieu account can reuse_session=N de tranh dung nham phien")

    tprint("\n" + "="*60)
    if errors or warnings:
        if errors:
            tprint(f"CONFIG INVALID - {len(errors)} loi, {len(warnings)} canh bao")
            for e in errors:
                tprint(f"  [ERROR] {e}")
        if warnings:
            tprint(f"CANH BAO - {len(warnings)} muc")
            for w in warnings:
                tprint(f"  [WARN]  {w}")
        tprint("="*60 + "\n")
        return len(errors) == 0
    else:
        tprint(f"CONFIG OK - {len(df_steps)} steps, {len(testcases)} testcases")
        tprint("="*60 + "\n")
        return True


# ── Run step ──────────────────────────────────────────────

def run_step(page: Page, step: pd.Series, value: str, tc_id: str = ""):
    action        = step["action"].lower()
    lt            = step["locator_type"]
    loc           = step["locator"]
    wait_sel      = step["wait_selector"]
    step_name     = step["step"]
    dropdown_sel  = step.get("dropdown_selector", "").strip()
    match_type    = step.get("match_type", "").strip()
    prefill_check = step.get("prefill_check", "").strip().upper() == "Y"

    display_value = "[REDACTED]" if value else ""
    tprint(
        f"    → {step_name} ({action})" + (f" = '{display_value}'" if display_value else ""),
        tc_id=tc_id
    )

    match action:
        case "fill":
            if value:
                if prefill_check:
                    try:
                        current = build_locator(page, lt, loc).input_value() or ""
                        if current.strip() == value.strip():
                            return
                    except Exception:
                        pass
                do_fill(page, lt, loc, value)
        case "force_fill":
            if value:
                if prefill_check:
                    try:
                        current = build_locator(page, lt, loc).input_value() or ""
                        if current.strip() == value.strip():
                            return
                    except Exception:
                        pass
                do_fill(page, lt, loc, value, force=True)
        case "fill_enter":
            if value:
                do_fill_enter(page, lt, loc, value)
        case "click":
            do_click(page, lt, loc, wait_sel, tc_id=tc_id)
        case "click_if_exists":
            do_click_if_exists(page, lt, loc, wait_sel, tc_id=tc_id)
        case "check":
            do_check(page, lt, loc)
        case "uncheck":
            build_locator(page, lt, loc).uncheck()
        case "upload":
            from runner_agent.upload import local_upload_path
            build_locator(page, lt, loc).set_input_files(local_upload_path(value))
        case "radio":
            if value:
                do_radio(page, lt, loc, value)
        case "select_antd":
            if value:
                ds = dropdown_sel if dropdown_sel else "visible"
                mt = match_type   if match_type   else "contains"
                do_select_antd(page, lt, loc, value, ds, mt)
        case "force_select_antd":
            if value:
                if prefill_check:
                    try:
                        current = build_locator(page, lt, loc).input_value() or ""
                        if current.strip() == value.strip():
                            return
                    except Exception:
                        pass
                ds = dropdown_sel if dropdown_sel else "not_hidden"
                mt = match_type   if match_type   else "exact"
                do_force_select_antd(page, lt, loc, value, ds, mt)
        case "select":
            if value:
                do_select(page, lt, loc, value)
        case "nth":
            if value:
                do_fill(page, lt, loc, value)
        case "form_item":
            if value:
                do_fill(page, "form_item", loc, value)
        case "fill_sequence":
            if value:
                do_fill_sequence(page, loc, value)
        case "expect" | "read_result" | "read_result_group" | "read_result_single":
            pass
        case "wait":
            do_wait(page, wait_sel, tc_id=tc_id)
        case _:
            tprint(f"    [WARN] action '{action}' chua duoc ho tro", tc_id=tc_id)


# ── Run repeat group ──────────────────────────────────────

def run_repeat_group(page: Page, group_name: str, group_steps: pd.DataFrame,
                     tc: dict, account: dict, tc_id: str = ""):
    value_steps = group_steps[group_steps["value_source"] == "testcase"]
    if value_steps.empty:
        return
    first_col = value_steps.iloc[0]["step"]
    raw_value = tc.get(first_col, "")
    if not raw_value:
        return
    items  = [v.strip() for v in raw_value.split(";")]
    n_rows = len(items)
    for i in range(n_rows):
        for _, step in group_steps.iterrows():
            action = step["action"].lower()
            loc    = step["locator"].replace("{i}", str(i))
            if step["value_source"] == "empty":
                do_click(page, step["locator_type"], loc, step["wait_selector"], tc_id=tc_id)
                continue
            col_name   = step["step"]
            raw        = tc.get(col_name, "")
            value_list = [v.strip() for v in raw.split(";")]
            value      = value_list[i] if i < len(value_list) else ""
            if not value:
                continue
            value = resolve_local_reference(value)
            match action:
                case "fill":
                    do_fill(page, step["locator_type"], loc, value)
                case "fill_enter":
                    do_fill_enter(page, step["locator_type"], loc, value)
                case "select_antd":
                    ds = step.get("dropdown_selector","") or "visible"
                    mt = step.get("match_type","")        or "contains"
                    do_select_antd(page, step["locator_type"], loc, value, ds, mt)
                case "force_select_antd":
                    ds = step.get("dropdown_selector","") or "not_hidden"
                    mt = step.get("match_type","")        or "exact"
                    do_force_select_antd(page, step["locator_type"], loc, value, ds, mt)
                case "select":
                    do_select(page, step["locator_type"], loc, value)


# ── Run screen ────────────────────────────────────────────

def resolve_value(step: pd.Series, tc: dict, account: dict) -> str:
    source = step["value_source"].lower().strip()
    match source:
        case "account":  value = account.get(step["step"], "")
        case "testcase": value = tc.get(step["step"], "")
        case "keyword":  value = tc.get("search_keyword", "")
        case _:          value = ""
    return resolve_local_reference(value)


def resolve_local_reference(value):
    """Resolve a whole-cell placeholder locally, including expected values."""
    match_ref = re.fullmatch(r"\$\{([A-Z][A-Z0-9_]*)\}", str(value))
    if match_ref:
        value = os.getenv(match_ref.group(1))
        if value is None:
            raise ValueError("Thieu bien local cho placeholder")
        _redactor.register(value)
    return value


def run_screen(page: Page, screen: str, df_steps: pd.DataFrame,
               tc: dict, account: dict, tc_id: str = ""):
    steps = df_steps[
        (df_steps["screen"] == screen) &
        (df_steps["active"].str.upper() == "Y")
    ]
    processed_groups = set()
    for _, step in steps.iterrows():
        group = step.get("group", "")
        if group:
            if group in processed_groups:
                continue
            group_steps = steps[steps["group"] == group]
            run_repeat_group(page, group, group_steps, tc, account, tc_id=tc_id)
            processed_groups.add(group)
        else:
            value = resolve_value(step, tc, account)
            run_step(page, step, value, tc_id=tc_id)


# ── Read & verify ─────────────────────────────────────────

def read_and_verify(page: Page, tc: dict, tc_id: str,
                     df_steps: pd.DataFrame,
                     mode: str = "group") -> tuple[bool, list]:
    action_filter = (
        ["read_result", "read_result_group"] if mode == "group"
        else ["read_result_single"]
    )
    read_steps = df_steps[
        (df_steps["screen"] == RESULT_SCREEN) &
        (df_steps["action"].str.lower().isin(action_filter)) &
        (df_steps["active"].str.upper() == "Y")
    ]
    if read_steps.empty:
        return True, []

    try:
        id_ho_so = (
            page.locator("span", has_text=ID_FIELD_LABEL)
            .locator("xpath=..")
            .locator(ID_FIELD_VALUE_SEL)
            .inner_text().strip()
        )
    except Exception:
        id_ho_so = ""

    if mode == "group":
        try:
            n_blocks = page.get_by_role(
                "button", name=re.compile(BLOCK_TRIGGER_SEL)
            ).count()
            if n_blocks == 0:
                n_blocks = 1
        except Exception:
            n_blocks = 1
        first_field = read_steps.iloc[0]["step"].replace("read_", "")
        n_expected  = 0
        while tc.get(f"expected_{first_field}_{n_expected}", ""):
            n_expected += 1
        if n_expected > 0 and n_blocks != n_expected:
            tprint(f"  [WARN] Block ky vong: {n_expected} | thuc te: {n_blocks}", tc_id=tc_id)
    else:
        n_blocks = 1

    all_pass    = True
    result_rows = []

    for i in range(n_blocks):
        if mode == "group" and n_blocks > 1:
            try:
                page.get_by_role(
                    "button", name=re.compile(BLOCK_TRIGGER_SEL)
                ).nth(i).click()
            except Exception as e:
                tprint(f"  [WARN] Khong mo duoc block {i}: {e}", tc_id=tc_id)

        label = f"Block {i}" if mode == "group" else "Single"
        tprint(f"\n  [RESULT {mode.upper()}] {label}", tc_id=tc_id)

        actual_values = {}
        for _, rs in read_steps.iterrows():
            field_name = rs["step"].replace("read_", "")
            actual_values[field_name] = read_result_field(page, rs, i, tc_id=tc_id)

        field_results = {}
        for fn, actual in actual_values.items():
            expected = (
                tc.get(f"expected_{fn}_{i}", "") if mode == "group"
                else tc.get(f"expected_{fn}", "")
            )
            expected = resolve_local_reference(expected)
            match_row = read_steps[read_steps["step"] == f"read_{fn}"]
            threshold = (
                match_row.iloc[0].get("match_type", "").strip()
                if not match_row.empty else ""
            ) or "exact"

            if not expected:
                tprint(f"    {fn:<20}: {actual} (khong kiem tra)", tc_id=tc_id)
                field_results[fn] = "N/A"
            else:
                passed, detail = _compare_values(expected, actual, threshold)
                result = "PASS" if passed else "FAIL"
                if not passed:
                    all_pass = False
                field_results[fn] = result
                tprint(
                    f"    {fn:<20}: ky vong={expected:<20} | "
                    f"thuc te={actual:<20} | {detail} → {result}",
                    tc_id=tc_id
                )

        checked_results = [v for v in field_results.values() if v != "N/A"]
        block_pass = all(v == "PASS" for v in checked_results)
        row = {
            "tc_id":    tc_id,
            "mo_ta":    tc.get("mo_ta", ""),
            "id_ho_so": id_ho_so,
            "overall":  ("PASS" if block_pass else "FAIL") if checked_results else "UNVERIFIED",
        }
        if mode == "group":
            row["block"] = i
        for fn, actual in actual_values.items():
            exp_key = f"expected_{fn}_{i}" if mode == "group" else f"expected_{fn}"
            row[f"exp_{fn}"]  = tc.get(exp_key, "")
            row[f"real_{fn}"] = actual
            row[f"pass_{fn}"] = field_results.get(fn, "N/A")

        result_rows.append(row)

    return all_pass, result_rows


# ── Export Excel ──────────────────────────────────────────

def export_results_simple(all_results: list, output_path: str):
    all_results = _redactor.value(all_results)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if not all_results:
        return

    all_keys = ["tc_id", "mo_ta", "id_ho_so", "overall"]
    seen = set(all_keys)
    field_keys = []
    for row in all_results:
        for k in row:
            if k not in seen:
                seen.add(k)
                field_keys.append(k)
    headers = all_keys + field_keys

    hdr_fill  = PatternFill("solid", start_color="4472C4")
    pass_fill = PatternFill("solid", start_color="E2EFDA")
    fail_fill = PatternFill("solid", start_color="FFCCCC")
    na_fill   = PatternFill("solid", start_color="F2F2F2")
    meta_fill = PatternFill("solid", start_color="F2F2F2")
    exp_fill  = PatternFill("solid", start_color="FFF2CC")
    real_fill = PatternFill("solid", start_color="D9E1F2")
    white_bold= Font(name="Arial", bold=True, color="FFFFFF", size=10)
    normal    = Font(name="Arial", size=10)
    center    = Alignment(horizontal="center", vertical="center")
    left      = Alignment(horizontal="left",   vertical="center")
    thin      = Side(style="thin", color="BFBFBF")
    bdr       = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook(); ws = wb.active; ws.title = "results"
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font=white_bold; cell.fill=hdr_fill
        cell.alignment=center; cell.border=bdr
        ws.column_dimensions[get_column_letter(c)].width = (
            12 if h in ("tc_id","overall","id_ho_so") else
            45 if h == "mo_ta" else
            10 if h.startswith("pass_") else 22
        )
    ws.row_dimensions[1].height = 28

    for r, row in enumerate(all_results, 2):
        for c, h in enumerate(headers, 1):
            val  = row.get(h, "")
            cell = ws.cell(row=r, column=c, value=val)
            cell.border=bdr; cell.font=normal; cell.alignment=left
            if h == "overall":
                cell.fill = pass_fill if val=="PASS" else fail_fill
                cell.alignment = center
            elif h in ("tc_id","mo_ta","id_ho_so"):
                cell.fill = meta_fill
            elif h.startswith("pass_"):
                cell.fill = (pass_fill if val=="PASS" else
                             fail_fill if val=="FAIL" else na_fill)
                cell.alignment = center
            elif h.startswith("exp_"):
                cell.fill = exp_fill
            elif h.startswith("real_"):
                cell.fill = real_fill

    ws.freeze_panes = "E2"
    wb.save(output_path)
    tprint(f"  Xuat ket qua: {output_path}")


def export_results_excel(all_results: list, output_path: str):
    all_results = _redactor.value(all_results)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if not all_results:
        return

    tc_map   = defaultdict(list)
    tc_order = []
    for row in all_results:
        tid = row.get("tc_id", "")
        if tid not in tc_map:
            tc_order.append(tid)
        tc_map[tid].append(row)

    n_blocks = max(len(b) for b in tc_map.values())

    META_EXCLUDE = {"id_ho_so", "tc_id", "mo_ta", "block", "overall"}
    seen_fields  = []
    for row in all_results:
        for k in row:
            if k.startswith("real_") and k[5:] not in seen_fields and k[5:] not in META_EXCLUDE:
                seen_fields.append(k[5:])
    FIELDS = seen_fields

    META    = ["tc_id", "mo_ta", "id_ho_so", "overall"]
    headers = META[:]
    widths  = [12, 45, 12, 10]
    for i in range(n_blocks):
        for f in FIELDS:
            headers += [f"b{i}_exp_{f}", f"b{i}_real_{f}", f"b{i}_pass_{f}"]
            widths  += [22, 22, 10]

    hdr_fill  = PatternFill("solid", start_color="4472C4")
    blk_fills = [
        PatternFill("solid", start_color="E8F4FD"),
        PatternFill("solid", start_color="FFF2CC"),
        PatternFill("solid", start_color="F0E8FD"),
        PatternFill("solid", start_color="E8FDE8"),
    ]
    pass_fill  = PatternFill("solid", start_color="E2EFDA")
    fail_fill  = PatternFill("solid", start_color="FFCCCC")
    na_fill    = PatternFill("solid", start_color="F2F2F2")
    meta_fill  = PatternFill("solid", start_color="F2F2F2")
    white_bold = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    normal     = Font(name="Arial", size=10)
    center     = Alignment(horizontal="center", vertical="center")
    left       = Alignment(horizontal="left",   vertical="center")
    thin       = Side(style="thin", color="BFBFBF")
    bdr        = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook(); ws = wb.active; ws.title = "results"
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font=white_bold; cell.fill=hdr_fill
        cell.alignment=center; cell.border=bdr
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 28

    for r, tid in enumerate(tc_order, 2):
        blocks  = tc_map[tid]
        mo_ta   = next((b.get("mo_ta","") for b in blocks if b.get("mo_ta")), "")
        id_hs   = blocks[0].get("id_ho_so","") if blocks else ""
        overall = ("FAIL" if any(b.get("overall") == "FAIL" for b in blocks)
                   else "UNVERIFIED" if any(b.get("overall") != "PASS" for b in blocks)
                   else "PASS")

        row_data = {"tc_id": tid, "mo_ta": mo_ta, "id_ho_so": id_hs, "overall": overall}
        for i, block in enumerate(blocks):
            for f in FIELDS:
                row_data[f"b{i}_exp_{f}"]  = block.get(f"exp_{f}",  "")
                row_data[f"b{i}_real_{f}"] = block.get(f"real_{f}", "")
                row_data[f"b{i}_pass_{f}"] = block.get(f"pass_{f}", "N/A")

        for c, h in enumerate(headers, 1):
            val  = row_data.get(h, "")
            cell = ws.cell(row=r, column=c, value=val)
            cell.border=bdr; cell.font=normal; cell.alignment=left
            if h == "overall":
                cell.fill = pass_fill if val=="PASS" else fail_fill
                cell.alignment = center
            elif h in META:
                cell.fill = meta_fill
            elif "_pass_" in h:
                cell.fill = (pass_fill if val=="PASS" else fail_fill if val=="FAIL" else na_fill)
                cell.alignment = center
            else:
                bidx = next((i for i in range(n_blocks) if h.startswith(f"b{i}_")), 0)
                cell.fill = blk_fills[bidx % len(blk_fills)]

    ws.freeze_panes = "E2"
    wb.save(output_path)
    tprint(f"  Xuat ket qua: {output_path}")


# ── Run testcase ──────────────────────────────────────────

def run_testcase(page: Page, tc: dict, df_steps: pd.DataFrame,
                 settings: dict, dirs: dict, log_path: Path,
                 is_first: bool = False):
    tc_id  = tc["tc_id"]
    mo_ta  = tc["mo_ta"]
    prefix = dirs["prefix"]

    # Ghi nhận ngay cả testcase lỗi/không có dòng kết quả.
    outcome = {"tc_id": tc_id, "status": "ERROR", "assertions": 0}
    _case_outcomes.append(outcome)

    tprint(f"\n{'='*60}", tc_id=tc_id)
    tprint(f"[START] {tc_id} - {mo_ta}", tc_id=tc_id)
    tprint(f"{'='*60}", tc_id=tc_id)

    try:
        account = get_account(tc["role_code"], tc["specialized_bank"])

        if is_first:
            tprint(f"  [Screen] {LOGIN_SCREEN}", tc_id=tc_id)
            page.goto(settings["url"])
            page.wait_for_load_state("networkidle")
            run_screen(page, LOGIN_SCREEN, df_steps, tc, account, tc_id=tc_id)
            page.wait_for_load_state("networkidle")
            tprint(f"  Login OK - {account['role_code']}/{account['specialized_bank']}", tc_id=tc_id)
        else:
            tprint("  [Nav] Ve trang chu...", tc_id=tc_id)
            try:
                page.locator(HOME_BUTTON_SELECTOR).click()
                page.wait_for_load_state("networkidle")
            except Exception as e:
                tprint(f"  [WARN] Khong click duoc nut trang chu: {e}", tc_id=tc_id)
                page.goto(settings["url"])
                page.wait_for_load_state("networkidle")

        screen_flow = settings.get("screen_flow", "").strip()
        screens     = [s.strip() for s in screen_flow.split(",") if s.strip()]

        all_pass    = True
        result_rows = []

        for screen in screens:
            if screen == LOGIN_SCREEN:
                continue

            steps_in = df_steps[
                (df_steps["screen"] == screen) &
                (df_steps["active"].str.upper() == "Y")
            ]

            if screen == RESULT_SCREEN:
                tprint(f"\n  [Screen] {screen}", tc_id=tc_id)
                # Cho trang ket qua render xong 1 lan duy nhat truoc khi doc
                if RESULT_WAIT_SELECTOR:
                    try:
                        page.wait_for_selector(
                            RESULT_WAIT_SELECTOR, state="attached", timeout=30000)
                    except Exception as e:
                        tprint(
                            f"  [WARN] Khong tim thay result selector "
                            f"'{RESULT_WAIT_SELECTOR}': {e}",
                            tc_id=tc_id
                        )
                group_steps = steps_in[
                    steps_in["action"].str.lower().isin(["read_result","read_result_group"])
                ]
                if not group_steps.empty:
                    passed, rows = read_and_verify(page, tc, tc_id, df_steps, mode="group")
                    if not passed:
                        all_pass = False
                    result_rows.extend(rows)
                single_steps = steps_in[
                    steps_in["action"].str.lower() == "read_result_single"
                ]
                if not single_steps.empty:
                    passed, rows = read_and_verify(page, tc, tc_id, df_steps, mode="single")
                    if not passed:
                        all_pass = False
                    result_rows.extend(rows)
                run_screen(page, screen, df_steps, tc, account, tc_id=tc_id)
            else:
                if steps_in.empty:
                    continue
                tprint(f"  [Screen] {screen}", tc_id=tc_id)
                run_screen(page, screen, df_steps, tc, account, tc_id=tc_id)

        checked = sum(v in ("PASS", "FAIL") for row in result_rows
                      for k, v in row.items() if k.startswith("pass_"))
        status = ("PASS" if all_pass else "FAIL") if checked else "UNVERIFIED"
        outcome.update(status=status, assertions=checked)
        write_log(log_path, tc_id, mo_ta, status)
        if status == "PASS":
            tprint(f"\n[PASS] {tc_id} - {mo_ta}", tc_id=tc_id)
        elif status == "FAIL":
            tprint(f"\n[FAIL] {tc_id} - {mo_ta} (ket qua khong khop expected)", tc_id=tc_id)
        else:
            tprint(f"\n[UNVERIFIED] {tc_id}", tc_id=tc_id)
        return result_rows

    except Exception as e:
        tprint(f"\n[FAIL] {tc_id} - {mo_ta}", tc_id=tc_id)
        # Playwright exception có thể chứa input/DOM; chỉ xuất loại lỗi.
        tprint(f"       {type(e).__name__}", tc_id=tc_id)
        write_log(log_path, tc_id, mo_ta, f"ERROR: {type(e).__name__}")
        if settings.get("screenshot_on_error","Y").upper() == "Y":
            try:
                ts   = datetime.now().strftime("%Y%m%d_%H%M%S")
                path = dirs["errors"] / f"error_{prefix}_{tc_id}_{ts}.png"
                # Fail-closed: chỉ chụp khi tất cả mask áp dụng được.
                masks = [page.locator("input, textarea, select, [contenteditable]")]
                masks += [page.get_by_text(v, exact=False) for v in _redactor._values]
                page.screenshot(path=str(path), mask=masks)
                tprint(f"       Screenshot: {path}", tc_id=tc_id)
            except Exception:
                pass
        return []


# ── Main ──────────────────────────────────────────────────

def main():
    args = parse_args()

    if args.config:
        config_path = Path(args.config)
    else:
        config_path = pick_config()

    if not config_path or not config_path.exists():
        tprint("Khong tim thay file config.")
        sys.exit(1)

    execute_config(config_path)


def execute_config(config_path, run_id=None, output_dir=None, env_path=None,
                   load_environment=True):
    """Chạy workbook hiện có; mỗi lần gọi phải ở process local riêng."""
    global _redactor, _case_outcomes
    import json
    import time
    config_path = Path(config_path)
    _redactor = Redactor()
    _case_outcomes = []
    started = time.monotonic()
    if load_environment:
        load_dotenv(dotenv_path=env_path or os.getenv("RUNNER_ENV_PATH") or get_base_dir() / ".env")

    settings  = load_settings(config_path)
    if output_dir is None:
        dirs = get_output_dirs(config_path)
    else:
        base = Path(output_dir)
        dirs = {"base": base, "logs": base, "results": base,
                "errors": base / "screenshots", "prefix": "run"}
        for key in ("logs", "results", "errors"):
            dirs[key].mkdir(parents=True, exist_ok=True)
    prefix    = dirs["prefix"]
    df_steps  = load_steps(config_path)
    testcases = load_testcases(config_path)
    for tc in testcases:
        for key, value in tc.items():
            if key not in ("tc_id", "active", "role_code", "specialized_bank"):
                _redactor.register(value)

    if not testcases:
        tprint("Khong co testcase nao active=Y")
        raise ValueError("Khong co testcase active")

    testcases = filter_testcases(testcases, settings)
    if not testcases:
        tprint("Khong co testcase nao thoa man dieu kien loc.")
        raise ValueError("Khong co testcase sau filter")

    global SPINNER_TIMEOUT, POPUP_TIMEOUT, POPUP_BUTTON_LOCATOR, DEBUG, BLOCK_TRIGGER_SEL
    global STEP_WRAPPER_SELECTOR, SPINNER_WRAPPER_CLASS, TICK_CLASS, STEP_LINK_XPATH
    global RESULT_WAIT_SELECTOR, ID_FIELD_LABEL, ID_FIELD_VALUE_SEL
    global RESULT_SCREEN, SPINNER_STROKE_COLOR, HOME_BUTTON_SELECTOR
    global LOGIN_SCREEN, REUSE_SESSION
    SPINNER_TIMEOUT       = int(settings.get("spinner_timeout",    "120000"))
    POPUP_TIMEOUT         = int(settings.get("popup_timeout",      "5000"))
    POPUP_BUTTON_LOCATOR  = settings.get("popup_button",           ".ant-modal-footer button.ant-btn-primary")
    STEP_WRAPPER_SELECTOR = settings.get("step_wrapper_selector",  ".z-20")
    SPINNER_WRAPPER_CLASS = settings.get("spinner_wrapper_class",  "bg-warning50")
    TICK_CLASS            = settings.get("tick_class",             "text-primaryColor")
    STEP_LINK_XPATH       = settings.get("step_link_xpath",        "following-sibling::div[contains(@class,'flex-col')][1]//a")
    DEBUG                 = settings.get("debug",                  "N").upper() == "Y"
    BLOCK_TRIGGER_SEL     = settings.get("block_trigger_selector", r"^Kết quả chấm điểm")
    RESULT_WAIT_SELECTOR  = settings.get("result_wait_selector",   "[id^='forms_SCORING_RESULT_elements_']")
    ID_FIELD_LABEL        = settings.get("id_field_label",         "ID hồ sơ")
    ID_FIELD_VALUE_SEL    = settings.get("id_field_value_selector","span.text-textPrimary")
    RESULT_SCREEN         = settings.get("result_screen",          "scoring_result")
    SPINNER_STROKE_COLOR  = settings.get("spinner_stroke_color",   "#F4600C")
    HOME_BUTTON_SELECTOR  = settings.get("home_button_selector",   "img.ant-image-img.cursor-pointer")
    LOGIN_SCREEN          = settings.get("login_screen",           "login")
    REUSE_SESSION         = settings.get("reuse_session",          "Y").upper() == "Y"

    if not validate_config(df_steps, testcases, settings):
        tprint("Dung lai do config co loi. Vui long kiem tra va sua lai.")
        raise ValueError("Config khong hop le")

    ts       = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = dirs["logs"] / ("run_log.txt" if run_id else f"run_log_{prefix}_{ts}.txt")
    with open(log_path, "w", encoding="utf-8") as f:
        f.write(f"{'='*70}\n")
        f.write(f"Run: {run_id or 'local'}\n")
        f.write(f"Run bat dau: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        f.write(f"{'='*70}\n")

    tprint(f"Bat dau chay {len(testcases)} testcase\n")
    tprint(f"Output: {dirs['base']}\n")

    browser_type     = settings.get("browser",           "chromium").strip().lower()
    headless         = settings.get("headless",          "N").upper() == "Y"
    slow_mo          = int(settings.get("slow_mo",       "0"))
    parallel_workers = int(settings.get("parallel_workers", "1"))

    all_results = []

    if REUSE_SESSION:
        # Dung chung 1 playwright + browser + context + page
        with sync_playwright() as p:
            if browser_type == "chrome":
                browser = p.chromium.launch(channel="chrome", headless=headless, slow_mo=slow_mo)
            elif browser_type == "msedge":
                browser = p.chromium.launch(channel="msedge", headless=headless, slow_mo=slow_mo)
            elif browser_type == "firefox":
                browser = p.firefox.launch(headless=headless, slow_mo=slow_mo)
            elif browser_type == "webkit":
                browser = p.webkit.launch(headless=headless, slow_mo=slow_mo)
            else:
                browser = p.chromium.launch(headless=headless, slow_mo=slow_mo)

            context = browser.new_context(
                user_agent=(
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/124.0.0.0 Safari/537.36"
                ),
            )
            context.add_init_script(
                "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"
            )
            page = context.new_page()
            page.set_default_timeout(int(settings.get("default_timeout","10000")))

            for idx, tc in enumerate(testcases):
                rows = run_testcase(
                    page, tc, df_steps, settings, dirs, log_path,
                    is_first=(idx == 0)
                )
                if rows:
                    all_results.extend(rows)

            page.close()
            context.close()
            browser.close()

    else:
        # Moi TC co playwright + browser rieng → an toan cho parallel
        def run_one_isolated(tc):
            """Moi worker tu quan ly toan bo playwright lifecycle."""
            with sync_playwright() as p:
                if browser_type == "chrome":
                    browser = p.chromium.launch(channel="chrome", headless=headless, slow_mo=slow_mo)
                elif browser_type == "msedge":
                    browser = p.chromium.launch(channel="msedge", headless=headless, slow_mo=slow_mo)
                elif browser_type == "firefox":
                    browser = p.firefox.launch(headless=headless, slow_mo=slow_mo)
                elif browser_type == "webkit":
                    browser = p.webkit.launch(headless=headless, slow_mo=slow_mo)
                else:
                    browser = p.chromium.launch(headless=headless, slow_mo=slow_mo)

                context = browser.new_context(
                    user_agent=(
                        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                        "AppleWebKit/537.36 (KHTML, like Gecko) "
                        "Chrome/124.0.0.0 Safari/537.36"
                    ),
                )
                context.add_init_script(
                    "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"
                )
                page = context.new_page()
                page.set_default_timeout(int(settings.get("default_timeout","10000")))

                try:
                    return run_testcase(
                        page, tc, df_steps, settings, dirs, log_path,
                        is_first=True
                    )
                finally:
                    page.close()
                    context.close()
                    browser.close()

        if parallel_workers <= 1:
            # Tuan tu
            for tc in testcases:
                rows = run_one_isolated(tc)
                if rows:
                    all_results.extend(rows)
        else:
            # Song song - moi worker co playwright rieng
            tprint(f"Chay song song {parallel_workers} workers...")
            results_map = {}
            with ThreadPoolExecutor(max_workers=parallel_workers) as executor:
                futures_map = {
                    executor.submit(run_one_isolated, tc): tc["tc_id"]
                    for tc in testcases
                }
                # as_completed nam trong with → xu ly ngay khi TC xong
                for future in as_completed(futures_map):
                    tc_id = futures_map[future]
                    try:
                        rows = future.result()
                        results_map[tc_id] = rows or []
                        tprint(f"  TC {tc_id} hoan thanh")
                    except Exception as e:
                        tprint(f"[ERROR] TC {tc_id} loi: {type(e).__name__}")
                        results_map[tc_id] = []

            # Giu thu tu goc
            for tc in testcases:
                all_results.extend(results_map.get(tc["tc_id"], []))

    if all_results:
        result_path = dirs["results"] / ("results.xlsx" if run_id else f"result_{prefix}_{ts}.xlsx")
        has_block = any("block" in r for r in all_results)
        if has_block:
            export_results_excel(all_results, str(result_path))
        else:
            export_results_simple(all_results, str(result_path))

    # A worker may fail before run_testcase (for example browser launch).
    # Every selected testcase must have an outcome before aggregating results.
    completed_ids = {c["tc_id"] for c in _case_outcomes}
    for tc in testcases:
        if tc["tc_id"] not in completed_ids:
            _case_outcomes.append({"tc_id": tc["tc_id"], "status": "ERROR"})
    counts = {s: sum(c["status"] == s for c in _case_outcomes)
              for s in ("PASS", "FAIL", "ERROR", "UNVERIFIED")}
    status = ("ERROR" if counts["ERROR"] else "FAILED" if counts["FAIL"]
              else "UNVERIFIED" if counts["UNVERIFIED"] else "PASSED")
    summary = {"run_id": run_id, "status": status, "passed": counts["PASS"],
               "failed": counts["FAIL"], "errors": counts["ERROR"],
               "unverified": counts["UNVERIFIED"], "duration": time.monotonic() - started,
               "cases": _redactor.value(_case_outcomes)}
    (dirs["base"] / "summary.json").write_text(json.dumps(summary), encoding="utf-8")
    errors = [{"tc_id": c["tc_id"], "overall": c["status"]}
              for c in _case_outcomes if c["status"] != "PASS"]
    if errors:
        export_results_simple(errors, str(dirs["base"] / "errors.xlsx"))
    return summary


if __name__ == "__main__":
    main()
