"""Local reviewed locator replacement; preserve source and export an inactive copy."""
from datetime import datetime, timezone
import hashlib
import io
from pathlib import Path
import re

from openpyxl import load_workbook

from .authoring import STRUCTURAL_SELECTOR
from .config import validate_workbook

DIRECT_ACTIONS = {"fill", "force_fill", "fill_enter", "click", "click_if_exists", "check",
                  "select", "select_antd", "force_select_antd", "nth"}


def require_supported_step(step):
    action = step.get("action", "").lower()
    if step.get("group"):
        raise ValueError("Repeated groups need manual repair")
    if action not in DIRECT_ACTIONS and not (
            action == "read_result_single" and step.get("read_method") in {"css_input", "css_disabled"}):
        raise ValueError("Action or read method requires manual repair")


def validate_candidate(selector):
    if not isinstance(selector, str) or len(selector) > 2000 or not STRUCTURAL_SELECTOR.fullmatch(selector):
        raise ValueError("Expected a structural CSS candidate")
    return selector


def candidate_matches_action(step, selector):
    tag = selector.split(" > ")[-1].split(":", 1)[0]
    action = step["action"].lower()
    if action in {"fill", "force_fill", "fill_enter", "nth"}:
        return tag in {"input", "textarea"}
    if action == "select":
        return tag == "select"
    if action == "check" or action == "read_result_single":
        return tag == "input" or (action == "read_result_single" and tag in {"textarea", "select"})
    return True


def check_candidate(page, step, selector):
    require_supported_step(step)
    validate_candidate(selector)
    if not candidate_matches_action(step, selector):
        return "INCOMPATIBLE_TAG"
    try:
        locator = page.locator(selector)
        if locator.count() != 1:
            return "NOT_UNIQUE"
        if not locator.is_visible():
            return "HIDDEN"
        if not locator.evaluate("element => element === window.__runnerRepairTarget"):
            return "STALE_SELECTION"
        return "READY_FOR_REVIEW"
    except Exception:
        return "CHECK_ERROR"


def export_repair(config, output, row, selector, source_sha256):
    """Caller must recheck live selection after human review, before exporting."""
    selector = validate_candidate(selector)
    source, target = Path(config), Path(output)
    if (source.suffix.lower() != ".xlsx" or source.is_symlink()
            or target.suffix.lower() != ".xlsx" or target.exists()
            or source.stat().st_size > 10 * 1024 * 1024):
        raise ValueError("Use a local workbook and a new .xlsx output path")
    content = source.read_bytes()
    if not re.fullmatch(r"[a-f0-9]{64}", source_sha256) or hashlib.sha256(content).hexdigest() != source_sha256:
        raise ValueError("Workbook changed since selection; inspect the new version")
    validate_workbook(content, require_settings=False)
    wb = load_workbook(io.BytesIO(content), keep_links=False)
    try:
        steps = wb["steps"]
        headers = [c.value for c in steps[1]]
        required = {"action", "locator", "locator_type", "active"}
        if not required <= set(headers) or len(headers) != len(set(headers)) or not 2 <= row <= steps.max_row:
            raise ValueError("Invalid step row or columns")
        step = {str(key): str(steps.cell(row, index + 1).value or "") for index, key in enumerate(headers)}
        require_supported_step(step)
        if not candidate_matches_action(step, selector):
            raise ValueError("Candidate tag incompatible with action")
        steps.cell(row, headers.index("locator_type") + 1, "css")
        steps.cell(row, headers.index("locator") + 1, selector)
        for sheet in (steps, wb["testcases"]):
            columns = [c.value for c in sheet[1]]
            if columns.count("active") != 1:
                raise ValueError("Missing or duplicate active column")
            for number in range(2, sheet.max_row + 1):
                sheet.cell(number, columns.index("active") + 1, "N")
        review = wb.create_sheet("repair_review")
        review.append(["item", "value"])
        for item, value in [
            ("source_sha256", source_sha256), ("step_row", row),
            ("changed_columns", "locator_type, locator; all steps/testcases set inactive"),
            ("locator_type", "css"), ("locator", selector),
            ("created_at", datetime.now(timezone.utc).isoformat()),
            ("review", "User confirmed local target; no testcase execution or assertion verification"),
            ("next", "Review remaining wait/dropdown/read settings; inspect again before activation"),
        ]:
            review.append([item, value])
        output_bytes = io.BytesIO()
        wb.save(output_bytes)
    finally:
        wb.close()
    result = output_bytes.getvalue()
    validate_workbook(result, require_settings=False)
    with target.open("xb") as file:
        file.write(result)
    return {"row": row, "source_sha256": source_sha256, "draft_sha256": hashlib.sha256(result).hexdigest()}
