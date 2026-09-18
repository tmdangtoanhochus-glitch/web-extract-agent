"""Local, read-only locator checks. No navigation, input values, DOM or screenshots."""
from collections import Counter
from datetime import datetime, timezone
from functools import lru_cache
import hashlib
import importlib.util
import io
from pathlib import Path
import re

from openpyxl import load_workbook

from .config import validate_workbook


@lru_cache(maxsize=1)
def runner_locator_builder():
    # Import definitions only: runner loads environment only in execute_config.
    source = Path(__file__).resolve().parents[1] / "docs" / "runner.py"
    spec = importlib.util.spec_from_file_location("inspector_runner", source)
    runner = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(runner)
    return runner.build_locator


def load_inspection_workbook(path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx" or path.is_symlink() or path.stat().st_size > 10 * 1024 * 1024:
        raise ValueError("Expected a local .xlsx workbook up to 10 MB")
    content = path.read_bytes()
    validate_workbook(content, require_settings=False)
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        rows = wb["steps"].iter_rows(values_only=True)
        columns = list(next(rows, ()))
        required = {"screen", "step", "action", "locator_type", "locator", "active", "wait_selector"}
        if not required <= set(columns) or len(columns) != len(set(columns)):
            raise ValueError("Invalid step columns")
        steps, screens = [], []
        for row_number, row in enumerate(rows, 2):
            if row_number > 2001:
                raise ValueError("Too many step rows")
            if not any(value is not None for value in row):
                continue
            data = {str(key): str(value).strip() if value is not None else ""
                    for key, value in zip(columns, row)}
            screen = data["screen"]
            if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_-]{0,79}", screen):
                raise ValueError("Screen must be an identifier")
            if screen not in screens:
                screens.append(screen)
            data["row"] = row_number
            data["screen_index"] = screens.index(screen) + 1
            steps.append(data)
        if not steps:
            raise ValueError("Workbook has no steps")
        return {"sha256": hashlib.sha256(content).hexdigest(), "screens": screens, "steps": steps}
    finally:
        wb.close()


def _check(page, kind, locator_type, selector, builder):
    result = {"kind": kind, "status": "MANUAL_REVIEW", "match_count": None}
    if not selector or selector in (":not(*)", "self::*[false()]", "RUNNER_REVIEW_REQUIRED"):
        result["status"] = "UNRESOLVED"
        return result
    if "{i}" in selector or len(selector) > 2000:
        return result
    if locator_type not in {"css", "xpath", "role", "text", "label", "filter", "form_item", "nth", "role_nth", ""}:
        return result
    try:
        locator = builder(page, locator_type, selector)
        count = locator.count()
        result["match_count"] = count
        if count == 0:
            result["status"] = "NOT_FOUND"
        elif count > 1:
            result["status"] = "AMBIGUOUS"
        else:
            result["status"] = "UNIQUE_VISIBLE" if locator.is_visible() else "HIDDEN"
    except Exception:
        # Selector parsing/browser errors can include DOM and literal values.
        result["status"] = "CHECK_ERROR"
    return result


def inspect_screen(page, workbook, screen_index, builder=None):
    if screen_index < 1 or screen_index > len(workbook["screens"]):
        raise ValueError("Unknown screen index")
    builder = builder or runner_locator_builder()
    results = []
    for step in workbook["steps"]:
        if step["screen_index"] != screen_index:
            continue
        action = step["action"].lower()
        checks = []
        if action == "wait":
            from .authoring import STRUCTURAL_SELECTOR
            selector = step.get("wait_selector", "")
            if len(selector) <= 2000 and STRUCTURAL_SELECTOR.fullmatch(selector):
                checks.append(_check(page, "wait", "css", selector, builder))
            else:
                checks.append({"kind": "wait", "status": "MANUAL_REVIEW", "match_count": None})
        elif action in {"read_result", "read_result_single", "read_result_group"}:
            if step.get("read_method") in {"css_input", "css_disabled"}:
                checks.append(_check(page, "target", "css", step["locator"], builder))
            else:
                checks.append({"kind": "target", "status": "MANUAL_REVIEW", "match_count": None})
        elif action in {"fill", "force_fill", "fill_enter", "click", "click_if_exists", "check",
                        "radio", "select", "select_antd", "force_select_antd", "nth", "form_item"}:
            kind = "form_item" if action == "form_item" else step["locator_type"].lower()
            checks.append(_check(page, "target", kind, step["locator"], builder))
        else:
            checks.append({"kind": "target", "status": "MANUAL_REVIEW", "match_count": None})
        # wait_selector contains Runner commands as well as selectors. Never execute it.
        if step["wait_selector"] and action != "wait":
            checks.append({"kind": "wait", "status": "MANUAL_REVIEW", "match_count": None})
        if action in {"select_antd", "force_select_antd"}:
            checks.append({"kind": "dropdown", "status": "MANUAL_REVIEW", "match_count": None})
        if action == "radio":
            checks.append({"kind": "value_dependent", "status": "MANUAL_REVIEW", "match_count": None})
        if step.get("group") or action in {"read_result", "read_result_group"}:
            checks.append({"kind": "group", "status": "MANUAL_REVIEW", "match_count": None})
        results.append({"row": step["row"], "checks": checks})
    return {"screen_index": screen_index, "observed_at": datetime.now(timezone.utc).isoformat(), "rows": results,
            "counts": dict(Counter(c["status"] for row in results for c in row["checks"]))}
