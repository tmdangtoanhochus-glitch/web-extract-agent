"""Merge generated steps into a user's config; settings change only with explicit approval."""
import io
from pathlib import Path
from openpyxl import load_workbook
from .config import validate_workbook
from src.runner.workbook_schema import STEP_COLUMNS, READ_ACTIONS, testcase_columns, header_result_blocks


def read_workbook(path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx" or path.is_symlink() or path.stat().st_size > 10 * 1024 * 1024:
        raise ValueError("Use a local .xlsx file up to 10 MB")
    return path.read_bytes()


def analyze(template_content, draft_content):
    validate_workbook(template_content)
    validate_workbook(draft_content, require_settings=False)
    template = load_workbook(io.BytesIO(template_content), keep_links=False)
    draft = load_workbook(io.BytesIO(draft_content), keep_links=False)
    try:
        setting_rows = list(template["settings"].values)
        if not setting_rows or list(setting_rows[0]) != ["key", "value"]:
            raise ValueError("Settings must have key/value columns")
        settings = {}
        for row in setting_rows[1:]:
            if row[0] is None:
                continue
            if row[0] in settings:
                raise ValueError("Duplicate settings key")
            settings[row[0]] = str(row[1] or "")
        raw = list(draft["steps"].values)
        if not raw or list(raw[0]) != STEP_COLUMNS or len(raw) > 2001:
            raise ValueError("Draft requires complete step columns")
        steps = [{key: str(value or "") for key, value in zip(STEP_COLUMNS, row)}
                 for row in raw[1:] if any(value is not None for value in row)]
        if not steps or any(step["active"] != "N" for step in steps):
            raise ValueError("Generated steps must be inactive")
        if any(any(value is not None for value in row) for row in list(draft["testcases"].values)[1:]):
            raise ValueError("Draft must contain testcase headers only; put user data in the template")
        header_result_blocks(steps, list(next(draft["testcases"].values)))
        screens = list(dict.fromkeys(step["screen"] for step in steps))
        proposals = []
        def propose(key, value, default, reason):
            old = settings.get(key, default)
            if old != value:
                proposals.append({"key": key, "before": old, "after": value, "reason": reason})
        old_flow = [s.strip() for s in settings.get("screen_flow", "").split(",") if s.strip()]
        if [s for s in old_flow if s in screens] != screens:
            propose("screen_flow", ",".join(screens), "",
                    "run_testcase/run_screen đọc screen_flow để chọn và sắp thứ tự màn hình; cấu hình hiện tại chưa bao phủ đúng các step mới.")
        login_screens = {s["screen"] for s in steps if s["value_source"] == "account"}
        result_screens = {s["screen"] for s in steps if s["action"] in READ_ACTIONS}
        if len(login_screens) > 1 or len(result_screens) > 1:
            raise ValueError("Multiple account/result screens require manual configuration")
        if login_screens:
            propose("login_screen", next(iter(login_screens)), "login",
                    "run_testcase thực hiện LOGIN_SCREEN trước flow; step dùng account nằm ở màn hình khác cấu hình hiện tại.")
        if result_screens:
            propose("result_screen", next(iter(result_screens)), "scoring_result",
                    "read_and_verify chỉ kiểm tra expected tại RESULT_SCREEN; cần khớp màn hình chứa read_result.")
        return steps, proposals
    finally:
        template.close()
        draft.close()


def merge(template_content, draft_content, approved_keys=()):
    steps, proposals = analyze(template_content, draft_content)
    approved = set(approved_keys)
    if not approved <= {p["key"] for p in proposals}:
        raise ValueError("Approval must refer to an explained proposal")
    wb = load_workbook(io.BytesIO(template_content), keep_links=False)
    try:
        index = wb.sheetnames.index("steps")
        del wb["steps"]
        sheet = wb.create_sheet("steps", index)
        sheet.append(STEP_COLUMNS)
        for step in steps:
            sheet.append([step[key] for key in STEP_COLUMNS])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        cases = wb["testcases"]
        headers = [cell.value for cell in cases[1]]
        if len(headers) != len(set(headers)):
            raise ValueError("Duplicate testcase headers")
        draft = load_workbook(io.BytesIO(draft_content), read_only=True, keep_links=False)
        try:
            draft_headers = list(next(draft["testcases"].values))
            # Keep compiler evidence/review alongside prepared steps, without replacing user notes.
            if "review" in draft.sheetnames:
                imported_review = wb.create_sheet("draft_review")
                for row in draft["review"].iter_rows(values_only=True):
                    imported_review.append(row)
                imported_review.freeze_panes = "A2"
        finally:
            draft.close()
        for key in draft_headers:
            if key not in headers:
                headers.append(key)
                cases.cell(1, len(headers), key)
        settings = wb["settings"]
        rows = {settings.cell(row, 1).value: row for row in range(2, settings.max_row + 1)}
        review = wb.create_sheet("preparation_review")
        review.append(["key", "before", "proposed", "reason", "decision"])
        for proposal in proposals:
            key = proposal["key"]
            if key in approved:
                if key in rows:
                    settings.cell(rows[key], 2, proposal["after"])
                else:
                    settings.append([key, proposal["after"]])
            review.append([key, proposal["before"], proposal["after"], proposal["reason"],
                           "APPROVED" if key in approved else "KEPT_ORIGINAL"])
        review.append(["testcases", "", "", "Existing user rows preserved; no testcase data generated. Enter new data yourself.", ""])
        review.append(["execution", "", "", "Steps stay inactive. Review locators, settings and user data before running.", ""])
        output = io.BytesIO()
        wb.save(output)
    finally:
        wb.close()
    validate_workbook(output.getvalue())
    return output.getvalue()
