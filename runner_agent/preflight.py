"""Static local checks only: no browser, environment loading or secret resolution."""
import io
from openpyxl import load_workbook
from .config import validate_workbook
from src.runner.workbook_schema import BASE_CASE_COLUMNS, READ_ACTIONS


def check(content):
    validate_workbook(content, require_settings=False)
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    issues, warnings = [], []
    try:
        def records(name, required):
            rows = wb[name].iter_rows(values_only=True)
            headers = list(next(rows, ()))
            if not required <= set(headers) or len(headers) != len(set(headers)):
                issues.append({"sheet": name, "code": "INVALID_HEADERS"})
                return []
            return [(number, {str(k): str(v).strip() if v is not None else "" for k, v in zip(headers, row)})
                    for number, row in enumerate(rows, 2) if any(v is not None for v in row)]
        if "settings" not in wb.sheetnames:
            issues.append({"sheet": "settings", "code": "MISSING_SETTINGS_USE_PREPARE"})
            settings = {}
        else:
            settings_rows = records("settings", {"key", "value"})
            settings = {record["key"]: record["value"] for _, record in settings_rows}
            if len(settings) != len(settings_rows):
                issues.append({"sheet": "settings", "code": "DUPLICATE_SETTING"})
        if not settings.get("url") or settings.get("url") == "https://example.invalid":
            issues.append({"sheet": "settings", "code": "MISSING_OR_PLACEHOLDER_URL"})
        steps = records("steps", {"screen", "step", "action", "locator", "locator_type", "value_source", "active", "wait_selector"})
        cases = records("testcases", set(BASE_CASE_COLUMNS))
        active_steps = [(row, s) for row, s in steps if s["active"].upper() == "Y"]
        active_cases = [(row, c) for row, c in cases if c["active"].upper() == "Y"]
        if not active_steps:
            issues.append({"sheet": "steps", "code": "NO_ACTIVE_STEPS"})
        if not active_cases:
            issues.append({"sheet": "testcases", "code": "ENTER_AND_ACTIVATE_USER_TESTCASES"})
        flow = {s.strip() for s in settings.get("screen_flow", "").split(",") if s.strip()}
        for row, step in active_steps:
            if step["screen"] not in flow:
                issues.append({"sheet": "steps", "row": row, "code": "SCREEN_NOT_IN_FLOW"})
            if step["action"] != "wait" and step["locator"] in {"", ":not(*)", "RUNNER_REVIEW_REQUIRED"}:
                issues.append({"sheet": "steps", "row": row, "code": "UNRESOLVED_LOCATOR"})
            if step["action"] == "wait" and step["wait_selector"] in {"", ":not(*)"}:
                issues.append({"sheet": "steps", "row": row, "code": "UNRESOLVED_WAIT"})
            if step["action"] in READ_ACTIONS and step["screen"] != settings.get("result_screen", "scoring_result"):
                issues.append({"sheet": "steps", "row": row, "code": "RESULT_SCREEN_MISMATCH"})
        seen = set()
        for row, case in active_cases:
            if not case["tc_id"] or case["tc_id"] in seen:
                issues.append({"sheet": "testcases", "row": row, "code": "MISSING_OR_DUPLICATE_ID"})
            seen.add(case["tc_id"])
        if not any(s["action"] in READ_ACTIONS for _, s in active_steps):
            warnings.append({"sheet": "steps", "code": "NO_ACTIVE_ASSERTIONS"})
        return {"scope": "static_only_no_secret_or_browser_check", "ready_for_local_review": not issues,
                "active_steps": len(active_steps), "active_testcases": len(active_cases),
                "issues": issues, "warnings": warnings}
    finally:
        wb.close()
