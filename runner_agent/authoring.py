"""Steps and testcase headers only. Settings and testcase data belong to users."""
import io
import re
from openpyxl import Workbook
from .config import validate_workbook
from src.runner.workbook_schema import STEP_COLUMNS, testcase_columns

STRUCTURAL_SELECTOR = re.compile(
    r"[a-z][a-z0-9-]*:nth-of-type\([1-9][0-9]{0,5}\)"
    r"(?: > [a-z][a-z0-9-]*:nth-of-type\([1-9][0-9]{0,5}\)){0,39}")


class Recording:
    def __init__(self, limit=1000):
        self.events, self.limit, self.dropped = [], limit, 0

    def accept(self, payload):
        if (not isinstance(payload, dict) or set(payload) != {"action", "locator"}
                or payload.get("action") not in ("click", "fill", "select")
                or not isinstance(payload.get("locator"), str)
                or len(payload["locator"]) > 2000
                or not STRUCTURAL_SELECTOR.fullmatch(payload["locator"])
                or len(self.events) >= self.limit):
            self.dropped += 1
            return False
        self.events.append(dict(payload))
        return True


def draft_workbook(recording):
    rows = []
    for index, event in enumerate(recording.events, 1):
        if not Recording().accept(event):
            raise ValueError("Invalid recording event")
        rows.append(dict(zip(STEP_COLUMNS, ["recorded", f"field_{index:04d}", event["action"],
            "css", event["locator"], "testcase" if event["action"] in {"fill", "select"} else "empty",
            "N", "", "", "", "N", "", ""])))
    return _workbook(rows, recording.dropped)


def planned_workbook(plan):
    from src.runner.workbook_schema import WorkbookPlan
    plan = WorkbookPlan.model_validate(plan.model_dump())
    return _workbook([step.model_dump() for step in plan.steps])


def _workbook(rows, dropped=0):
    wb = Workbook()
    steps = wb.active
    steps.title = "steps"
    steps.append(STEP_COLUMNS)
    for row in rows:
        steps.append([row[key] for key in STEP_COLUMNS])
    cases = wb.create_sheet("testcases")
    cases.append(testcase_columns(rows))
    notes = wb.create_sheet("review")
    notes.append(["item", "instruction"])
    for item in [
        ("scope", "Steps and testcase HEADERS only. Enter all testcase rows yourself."),
        ("settings", "No settings generated. Use prepare_runner.py with your existing config; changes require explanation and approval."),
        ("locators", "Review locators/read_method/wait/dropdown locally before activating any step."),
        ("credentials", "Use local account references/placeholders; never put real credentials in the workbook."),
        ("assertions", "Enter expected data yourself; read steps do not establish business expectations."),
        ("dropped_events", str(dropped)),
    ]:
        notes.append(item)
    return _save_workbook(wb)


def _save_workbook(wb):
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    content = output.getvalue()
    validate_workbook(content, require_settings=False)
    return content
