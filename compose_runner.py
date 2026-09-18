"""Combine reviewed, header-only drafts in explicit order without generating data."""
import argparse
import io
from pathlib import Path

from openpyxl import load_workbook

from runner_agent.authoring import _workbook
from runner_agent.config import validate_workbook
from runner_agent.preparation import read_workbook
from src.runner.workbook_schema import STEP_COLUMNS, BASE_CASE_COLUMNS, READ_ACTIONS, header_result_blocks


def compose(drafts, output):
    target = Path(output)
    if target.suffix.lower() != ".xlsx" or target.exists() or not 1 <= len(drafts) <= 20:
        raise ValueError("Use 1..20 drafts in execution order and a new .xlsx output")
    rows, names, screens, result_screens = [], set(), [], set()
    source_reviews = []
    result_blocks = 1
    group_owner = {}
    sources = [(Path(path), read_workbook(path)) for path in drafts]
    for source_index, (_, content) in enumerate(sources, 1):
        validate_workbook(content, require_settings=False)
        wb = load_workbook(io.BytesIO(content), read_only=True, keep_links=False)
        try:
            if "review" in wb.sheetnames:
                for note in wb["review"].iter_rows(min_row=2, values_only=True):
                    if any(value is not None for value in note):
                        source_reviews.append((f"draft_{source_index}", " | ".join(str(value or "") for value in note)))
            if "settings" in wb.sheetnames:
                raise ValueError("Use generated drafts without settings; user config belongs in prepare_runner")
            raw = list(wb["steps"].values)
            if not raw or list(raw[0]) != STEP_COLUMNS:
                raise ValueError("Draft requires all step columns")
            current = [{key: str(value or "") for key, value in zip(STEP_COLUMNS, row)}
                       for row in raw[1:] if any(value is not None for value in row)]
            cases = list(wb["testcases"].values)
            if (not current or not cases or
                    any(any(value is not None for value in row) for row in cases[1:])):
                raise ValueError("Draft must have matching testcase headers and no testcase rows")
            result_blocks = max(result_blocks, header_result_blocks(current, list(cases[0])))
            for group in {s["group"] for s in current if s["group"]}:
                if group in group_owner:
                    raise ValueError("Repeat group names must be unique across drafts")
                group_owner[group] = True
            for step in current:
                if step["active"] != "N" or not step["step"] or step["step"] in names:
                    raise ValueError("Keep steps inactive and names unique across drafts")
                names.add(step["step"])
                screen = step["screen"]
                if not screen:
                    raise ValueError("Every step requires a screen")
                if not screens or screens[-1] != screen:
                    if screen in screens:
                        raise ValueError("A screen must stay contiguous; executor groups steps by screen")
                    screens.append(screen)
                if step["action"] in READ_ACTIONS:
                    result_screens.add(screen)
                rows.append(step)
        finally:
            wb.close()
    data = {s["step"] for s in rows if s["value_source"] == "testcase"}
    expected = {"expected_" + s["step"].removeprefix("read_") +
                ("" if s["action"] == "read_result_single" else "_0")
                for s in rows if s["action"] in READ_ACTIONS}
    if len(rows) > 2000 or len(result_screens) > 1 or data & (set(BASE_CASE_COLUMNS) | expected):
        raise ValueError("Draft exceeds executor limits or has conflicting testcase columns")
    result = _workbook(rows, result_blocks=result_blocks)
    if source_reviews:
        combined = load_workbook(io.BytesIO(result), keep_links=False)
        try:
            for note in source_reviews:
                combined["review"].append(note)
            buffer = io.BytesIO()
            combined.save(buffer)
            result = buffer.getvalue()
        finally:
            combined.close()
        validate_workbook(result, require_settings=False)
    if any(read_workbook(path) != content for path, content in sources):
        raise ValueError("A draft changed while composing")
    with target.open("xb") as file:
        file.write(result)
    return {"drafts": len(sources), "steps": len(rows), "screens": len(screens)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--draft", required=True, action="append", help="Repeat in execution order")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    try:
        result = compose(args.draft, args.output)
        print(f"Combined {result['drafts']} drafts: {result['steps']} inactive steps, {result['screens']} screens. No testcase data/settings generated.")
    except Exception as error:
        raise SystemExit(f"Composition stopped ({type(error).__name__}); source files unchanged.") from None
